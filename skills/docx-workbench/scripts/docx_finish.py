#!/usr/bin/env python3
"""Add the finishing parts of a report to an existing .docx (typically pandoc output):
cover page, table of contents with real page numbers, page-number footer,
running header, paper size and margins.

    python docx_finish.py body.docx -o report.docx --a4 --cover --toc --page-numbers \
        [--title "..."] [--subtitle "..."] [--author "..."] [--date "3 October 2026"] [--report-no "GL-2026-014"] \
        [--header-text "Slope Stability – Lot 12"] [--toc-levels 3] [--margins-cm 2.5]

Cover: uses --title/--author/--date, else the document's Title/Author/Date paragraphs
       (pandoc puts YAML front matter there), else core properties. The cover gets
       no header/footer and counts as page 1.
TOC:   a real Word TOC field ("Contents"), pre-filled with entries and page numbers
       (found by rendering with LibreOffice), so it is correct on opening; Word can
       still refresh it (right-click → Update Field).
"""
import argparse
import datetime as dt
import os
import re
import shutil
import subprocess
import sys
import tempfile

sys.dont_write_bytecode = True
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from _docx_common import add_field, ensure_paragraph_style, set_style_id  # noqa: E402

import docx  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt  # noqa: E402
from docx.text.paragraph import Paragraph  # noqa: E402


def style_id(p):
    ps = p._p.find(f"{qn('w:pPr')}/{qn('w:pStyle')}")
    return ps.get(qn("w:val")) if ps is not None else ""


def heading_level(p):
    sid = style_id(p).lower().replace(" ", "")
    m = re.match(r"heading(\d)$", sid)
    if m:
        return int(m.group(1))
    name = (p.style.name if p.style is not None else "").lower()
    m = re.match(r"heading (\d)$", name)
    return int(m.group(1)) if m else None


def new_paragraph_before(anchor_p, text="", style=None):
    el = OxmlElement("w:p")
    anchor_p._p.addprevious(el)
    p = Paragraph(el, anchor_p._parent)
    if style:
        set_style_id(p, style)
    if text:
        p.add_run(text)
    return p


def text_width_twips(section):
    return int((section.page_width - section.left_margin - section.right_margin) / 635)  # EMU -> twips


def build_cover(doc, a):
    body_ps = doc.paragraphs
    meta = {}
    used = []
    for p in body_ps[:6]:
        sid = style_id(p)
        if sid in ("Title", "Subtitle", "Author", "Date", "Abstract") and p.text.strip():
            meta[sid] = p.text.strip()
            used.append(p)
        elif p.text.strip():
            break
    cp = doc.core_properties
    title = a.title or meta.get("Title") or cp.title or "Report"
    subtitle = a.subtitle or meta.get("Subtitle")
    author = a.author or meta.get("Author") or cp.author
    date = a.date or meta.get("Date") or dt.date.today().strftime("%-d %B %Y")
    for p in used:  # rebuild the cover from scratch for consistent spacing
        p._p.getparent().remove(p._p)
    first = doc.paragraphs[0]
    tp = new_paragraph_before(first, title, "Title")
    tp.paragraph_format.space_before = Pt(160)
    tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if subtitle:
        sp = new_paragraph_before(first, subtitle, "Subtitle")
        sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    gap = new_paragraph_before(first)
    gap.paragraph_format.space_before = Pt(60)
    for line in filter(None, [author, a.report_no and f"Report {a.report_no}", date]):
        lp = new_paragraph_before(first, line)
        lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    brk = new_paragraph_before(first)
    brk.add_run().add_break(WD_BREAK.PAGE)
    if not cp.title:
        cp.title = title
    if author and not cp.author:
        cp.author = author
    return title


def add_bookmark(p, name, bid):
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bid))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bid))
    ppr = p._p.find(qn("w:pPr"))
    if ppr is not None:
        ppr.addnext(start)
    else:
        p._p.insert(0, start)
    p._p.append(end)


def build_toc(doc, a, after_cover):
    sec = doc.sections[0]
    tw = text_width_twips(sec)
    ensure_paragraph_style(doc, "TOCHeading", "TOC Heading", based_on="Heading1" if any(s.style_id == "Heading1" for s in doc.styles) else "Normal",
                           font_size=16, bold=True, space_after=12)
    for lvl in range(1, 4):
        ensure_paragraph_style(doc, f"TOC{lvl}", f"toc {lvl}", tab_right_twips=tw, indent_twips=(lvl - 1) * 220,
                               space_after=3 if lvl > 1 else 4, bold=(lvl == 1))
    heads = [(p, heading_level(p)) for p in doc.paragraphs]
    heads = [(p, lvl) for p, lvl in heads if lvl and lvl <= a.toc_levels and p.text.strip()]
    if not heads:
        print("No Heading-styled paragraphs found; TOC skipped", file=sys.stderr)
        return []
    entries = []
    for i, (p, lvl) in enumerate(heads):
        name = f"_Toc{900000 + i}"
        add_bookmark(p, name, 900000 + i)
        entries.append((p.text.strip(), lvl, name))

    # insertion point: first paragraph after the cover (or the start of the document)
    anchor = doc.paragraphs[after_cover] if after_cover < len(doc.paragraphs) else doc.paragraphs[-1]
    hp = new_paragraph_before(anchor, "Contents", "TOCHeading")
    toc_ps = []
    for k, (text, lvl, name) in enumerate(entries):
        ep = new_paragraph_before(anchor, style=f"TOC{lvl}")
        if k == 0:
            add_field_begin(ep, rf'TOC \o "1-{a.toc_levels}" \h \z \u')
        r = ep.add_run(text)
        ep.add_run("\t")
        add_field(ep, f"PAGEREF {name} \\h", "#")
        toc_ps.append(ep)
    end_r = toc_ps[-1].add_run()
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "end")
    end_r._r.append(fc)
    brk = new_paragraph_before(anchor)
    brk.add_run().add_break(WD_BREAK.PAGE)
    del r, hp
    return entries


def add_field_begin(p, instr):
    for kind, text in (("begin", None), (None, instr), ("separate", None)):
        r = p.add_run()
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            el.text = f" {text} "
        r._r.append(el)


def fill_toc_numbers(path, entries):
    """Render with LibreOffice, find each heading's page, write numbers into the cached PAGEREF results."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        print("LibreOffice not found: TOC page numbers left as '#'; Word fills them on Update Field", file=sys.stderr)
        return False
    try:
        import pypdfium2 as pdfium
    except ImportError:
        print("pypdfium2 missing: TOC page numbers left as '#'", file=sys.stderr)
        return False
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run([soffice, f"-env:UserInstallation=file://{tmp}/p", "--headless", "--convert-to", "pdf", "--outdir", tmp, path],
                       capture_output=True, timeout=240)
        pdf = os.path.join(tmp, os.path.splitext(os.path.basename(path))[0] + ".pdf")
        if not os.path.exists(pdf):
            print("LibreOffice render failed: TOC numbers left as '#'", file=sys.stderr)
            return False
        doc = pdfium.PdfDocument(pdf)
        norm = lambda s: re.sub(r"\s+", " ", s).strip().lower()  # noqa: E731
        pages = [norm(doc[i].get_textpage().get_text_range()) for i in range(len(doc))]
        doc.close()
    # skip pages up to and including the TOC itself (it repeats every heading)
    toc_page = next((i for i, t in enumerate(pages) if "contents" in t[:200]), 0)
    found, start = {}, toc_page + 1
    for text, lvl, name in entries:
        key = norm(text)[:60]
        for i in range(start, len(pages)):
            if key in pages[i]:
                found[name] = i + 1
                start = i
                break
    d = docx.Document(path)
    for p in d.paragraphs:
        instr = "".join(t.text or "" for t in p._p.iter(qn("w:instrText")))
        m = re.search(r"PAGEREF (_Toc\d+)", instr)
        if not m or m.group(1) not in found:
            continue
        # cached result = text runs between 'separate' and 'end' of the PAGEREF field
        state = None
        for r in p._p.iter(qn("w:r")):
            fc = r.find(qn("w:fldChar"))
            it = r.find(qn("w:instrText"))
            if it is not None and "PAGEREF" in (it.text or ""):
                state = "instr"
            elif fc is not None and fc.get(qn("w:fldCharType")) == "separate" and state == "instr":
                state = "result"
            elif fc is not None and fc.get(qn("w:fldCharType")) == "end" and state == "result":
                break
            elif state == "result":
                t = r.find(qn("w:t"))
                if t is not None:
                    t.text = str(found[m.group(1)])
                    state = "done"
    d.save(path)
    missing = [t for t, _, n in entries if n not in found]
    if missing:
        print(f"Could not locate page for: {missing} (left as '#')", file=sys.stderr)
    return True


def add_page_numbers(doc, a):
    ensure_paragraph_style(doc, "Footer", "footer")
    ensure_paragraph_style(doc, "Header", "header")
    sec = doc.sections[0]
    fp = sec.footer.paragraphs[0] if sec.footer.paragraphs else sec.footer.add_paragraph()
    for r in list(fp.runs):
        r._r.getparent().remove(r._r)
    set_style_id(fp, "Footer")
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    parts = re.split(r"(\{PAGE\}|\{NUMPAGES\})", a.page_format)
    for part in parts:
        if part == "{PAGE}":
            add_field(fp, "PAGE", "1")
        elif part == "{NUMPAGES}":
            add_field(fp, "NUMPAGES", "1")
        elif part:
            fp.add_run(part)
    if a.header_text:
        hp = sec.header.paragraphs[0] if sec.header.paragraphs else sec.header.add_paragraph()
        hp.text = a.header_text
        set_style_id(hp, "Header")
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("docx")
    ap.add_argument("-o", "--output", required=True)
    ap.add_argument("--a4", action="store_true")
    ap.add_argument("--letter", action="store_true")
    ap.add_argument("--margins-cm", type=float)
    ap.add_argument("--cover", action="store_true")
    ap.add_argument("--title"); ap.add_argument("--subtitle"); ap.add_argument("--author")
    ap.add_argument("--date"); ap.add_argument("--report-no")
    ap.add_argument("--toc", action="store_true")
    ap.add_argument("--toc-levels", type=int, default=3)
    ap.add_argument("--page-numbers", action="store_true")
    ap.add_argument("--page-format", default="Page {PAGE} of {NUMPAGES}")
    ap.add_argument("--header-text")
    a = ap.parse_args()
    if os.path.abspath(a.output) == os.path.abspath(a.docx):
        sys.exit("Refusing to overwrite the input; choose a different -o")

    doc = docx.Document(a.docx)
    for s in doc.sections:
        if a.a4 or a.letter:
            w, h = (Cm(21.0), Cm(29.7)) if a.a4 else (Cm(21.59), Cm(27.94))
            landscape = s.page_width and s.page_height and s.page_width > s.page_height
            s.page_width, s.page_height = (h, w) if landscape else (w, h)
        if a.margins_cm is not None:
            s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(a.margins_cm)

    after_cover = 0
    if a.cover:
        build_cover(doc, a)
        after_cover = next(i for i, p in enumerate(doc.paragraphs) if p._p.find(f".//{qn('w:br')}[@{qn('w:type')}='page']") is not None) + 1
        s0 = doc.sections[0]
        s0.different_first_page_header_footer = True
        for hf in (s0.first_page_header, s0.first_page_footer):  # explicit empty cover header/footer
            hf.is_linked_to_previous = False
            for para in hf.paragraphs:
                for r in list(para.runs):
                    r._r.getparent().remove(r._r)
    if a.page_numbers or a.header_text:
        add_page_numbers(doc, a)
    entries = build_toc(doc, a, after_cover) if a.toc else []
    doc.save(a.output)
    if entries:
        fill_toc_numbers(a.output, entries)
    done = [x for x, on in (("cover", a.cover), (f"TOC ({len(entries)} entries)", a.toc), ("page numbers", a.page_numbers),
                            ("header", bool(a.header_text)), ("A4" if a.a4 else "Letter", a.a4 or a.letter)) if on]
    print(f"Wrote {a.output}: " + ", ".join(done), file=sys.stderr)
    print("Check: python docx_convert.py " + a.output + " --to png --pages 1-3", file=sys.stderr)


if __name__ == "__main__":
    main()
