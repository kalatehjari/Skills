#!/usr/bin/env python3
"""Fill a Word template that contains {{placeholders}} from a JSON file.

Template rules
  {{name}}            replaced by data["name"] (works even if Word split it across runs)
  {{client.address}}  dotted paths reach into nested objects
  Table rows          a row containing {{items.<field>}} is repeated once per entry
                      of the list data["items"]; the original row is the pattern.
  Images              a value like {"image": "logo.png", "width_cm": 4} replaces the
                      placeholder with that picture (put it alone in its paragraph)
Placeholders work in body, tables, headers and footers.

Usage:
    python docx_fill_template.py template.docx data.json -o filled.docx [--strict]
--strict fails if any placeholder has no value (default: leave it and warn).
Batch: call once per record (see references/editing-and-review.md).
"""
import argparse
import copy
import json
import os
import re
import sys

sys.dont_write_bytecode = True
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from _docx_common import iter_paragraphs, replace_in_paragraph, text_runs  # noqa: E402

import docx  # noqa: E402
from docx.shared import Cm  # noqa: E402

PH = re.compile(r"\{\{\s*([A-Za-z0-9_.\-]+)\s*\}\}")
MISSING = object()


def lookup(data, path):
    cur = data
    for part in path.split("."):
        if isinstance(cur, dict) and part in cur:
            cur = cur[part]
        elif isinstance(cur, list) and part.isdigit() and int(part) < len(cur):
            cur = cur[int(part)]
        else:
            return MISSING
    return cur


def fmt(v):
    if isinstance(v, bool):
        return "Yes" if v else "No"
    if v is None:
        return ""
    return str(v)  # numbers keep their JSON form (8.0 stays "8.0"); pass strings for custom formatting


def expand_table_rows(doc, data):
    """Repeat rows that reference a list (e.g. {{bh.depth}} with data['bh'] = [...])."""
    count = 0
    for table in iter_tables(doc):
        for row in list(table.rows):
            keys = {m.group(1).split(".")[0] for c in row.cells for m in PH.finditer(c.text)}
            lists = [k for k in keys if isinstance(data.get(k), list)]
            if not lists:
                continue
            key = lists[0]
            items = data[key]
            tr = row._tr
            anchor = tr
            for idx, item in enumerate(items):
                new_tr = copy.deepcopy(tr)
                anchor.addnext(new_tr)
                anchor = new_tr
                new_row = docx.table._Row(new_tr, table)
                local = {**data, key: item}
                seen = set()
                for cell in new_row.cells:
                    if cell._tc in seen:
                        continue
                    seen.add(cell._tc)
                    for p in cell.paragraphs:
                        CELLS[0] += replace_in_paragraph(p, PH, lambda m: render(m, local, idx))
            tr.getparent().remove(tr)
            count += len(items)
    return count


def iter_tables(doc):
    def walk(container):
        for t in container.tables:
            yield t
            for row in t.rows:
                for cell in row.cells:
                    yield from walk(cell)
    yield from walk(doc)
    for s in doc.sections:
        for hf in (s.header, s.footer):
            if not hf.is_linked_to_previous:
                yield from walk(hf)


UNFILLED = []
CELLS = [0]


def render(m, data, idx=None):
    key = m.group(1)
    if key in ("#", "index") and idx is not None:
        return str(idx + 1)
    v = lookup(data, key)
    if v is MISSING:
        UNFILLED.append(key)
        return m.group(0)
    if isinstance(v, dict) and "image" in v:
        return m.group(0)  # images are handled separately
    return fmt(v)


def place_images(doc, data, base_dir):
    n = 0
    for p in iter_paragraphs(doc):
        for m in list(PH.finditer(p.text)):
            v = lookup(data, m.group(1))
            if isinstance(v, dict) and "image" in v:
                path = v["image"] if os.path.isabs(v["image"]) else os.path.join(base_dir, v["image"])
                if not os.path.exists(path):
                    sys.exit(f"Image not found for {{{{{m.group(1)}}}}}: {path}")
                replace_in_paragraph(p, re.compile(re.escape(m.group(0))), "")
                runs = text_runs(p)
                run = runs[0] if runs else p.add_run()
                kw = {}
                if "width_cm" in v:
                    kw["width"] = Cm(float(v["width_cm"]))
                if "height_cm" in v:
                    kw["height"] = Cm(float(v["height_cm"]))
                run.add_picture(path, **kw)
                n += 1
    return n


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("template")
    ap.add_argument("data")
    ap.add_argument("-o", "--output", required=True)
    ap.add_argument("--strict", action="store_true")
    a = ap.parse_args()
    if os.path.abspath(a.output) == os.path.abspath(a.template):
        sys.exit("Refusing to overwrite the template; choose a different -o")
    with open(a.data, encoding="utf-8") as fh:
        data = json.load(fh)

    d = docx.Document(a.template)
    rows = expand_table_rows(d, data)
    imgs = place_images(d, data, os.path.dirname(os.path.abspath(a.data)))
    n = 0
    for p in iter_paragraphs(d):
        n += replace_in_paragraph(p, PH, lambda m: render(m, data))
    missing = sorted(set(UNFILLED))
    if missing and a.strict:
        sys.exit(f"No value for: {missing}")
    d.save(a.output)
    print(f"Wrote {a.output}: {n} placeholder(s) filled in text/headers, {rows} table row(s) generated "
          f"({CELLS[0]} cell placeholder(s)), {imgs} image(s) placed", file=sys.stderr)
    if missing:
        print(f"Warning: left unfilled (no data): {missing}", file=sys.stderr)


if __name__ == "__main__":
    main()
