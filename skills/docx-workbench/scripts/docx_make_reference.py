#!/usr/bin/env python3
"""Create a styles template ("reference doc") so Markdown -> Word conversions
come out in a consistent house style.

    python docx_make_reference.py -o house.docx --font "Calibri" --size 11 \\
        --heading-font "Calibri Light" --heading-color 1F4E79 --margins-cm 2.5 --a4
    python docx_convert.py notes.md --to docx --reference-doc house.docx

Starts from pandoc's default reference.docx (so every style pandoc uses exists),
or from --base your_template.docx to keep an existing look and only adjust it.
"""
import argparse
import os
import subprocess
import sys
import tempfile

try:
    import docx
    import docx.opc.exceptions
    from docx.enum.text import WD_LINE_SPACING  # noqa: F401
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ImportError:
    sys.exit("python-docx is required: pip install python-docx")


def set_font(style, name=None, size=None, color=None, bold=None):
    f = style.font
    if name:
        f.name = name
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = rpr.makeelement(qn("w:rFonts"), {})
            rpr.append(rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(qn(attr), name)
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            if rfonts.get(qn(attr)) is not None:
                del rfonts.attrib[qn(attr)]
    if size:
        f.size = Pt(size)
    if color:
        f.color.rgb = RGBColor.from_string(color.upper())
    if bold is not None:
        f.bold = bold


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("-o", "--output", required=True)
    ap.add_argument("--base", help="start from this .docx instead of pandoc's default")
    ap.add_argument("--font", default="Calibri")
    ap.add_argument("--size", type=float, default=11)
    ap.add_argument("--heading-font")
    ap.add_argument("--heading-color", default="1F3864", help="hex RRGGBB (default 1F3864 = dark blue)")
    ap.add_argument("--line-spacing", type=float, default=1.15)
    ap.add_argument("--space-after", type=float, default=6, help="points after body paragraphs")
    ap.add_argument("--margins-cm", type=float, default=2.5)
    ap.add_argument("--a4", action="store_true", help="A4 paper (default keeps the base's size)")
    a = ap.parse_args()

    if a.base:
        d = docx.Document(a.base)
    else:
        with tempfile.TemporaryDirectory() as t:
            ref = os.path.join(t, "reference.docx")
            try:
                subprocess.run(["pandoc", "-o", ref, "--print-default-data-file", "reference.docx"], check=True)
                d = docx.Document(ref)
            except (FileNotFoundError, subprocess.CalledProcessError, docx.opc.exceptions.PackageNotFoundError):
                print("pandoc not available; starting from python-docx's default template", file=sys.stderr)
                d = docx.Document()

    # look styles up by displayed name (pandoc's template capitalises built-in names,
    # which python-docx's styles[...] lookup does not expect)
    by_name = {st.name: st for st in d.styles}
    hfont = a.heading_font or a.font
    for name in ("Normal", "Body Text", "First Paragraph", "Compact", "Block Text"):
        if name in by_name:
            st = by_name[name]
            set_font(st, a.font, a.size if name != "Compact" else None)
            pf = st.paragraph_format
            pf.line_spacing = a.line_spacing
            pf.space_after = Pt(a.space_after)
    sizes = {"Title": a.size * 2.2, "Subtitle": a.size * 1.4, "Heading 1": a.size * 1.6, "Heading 2": a.size * 1.3,
             "Heading 3": a.size * 1.15, "Heading 4": a.size}
    for name, sz in sizes.items():
        if name in by_name:
            set_font(by_name[name], hfont, round(sz), a.heading_color, bold=name.startswith("Heading"))
    for name in ("Caption", "Image Caption", "Table Caption"):
        if name in by_name:
            set_font(by_name[name], a.font, a.size - 1, "404040")
            by_name[name].font.italic = True

    # document defaults: anything without an explicit font (Author, Date, footnotes...) uses these
    rpr_default = d.styles.element.find(f"{qn('w:docDefaults')}/{qn('w:rPrDefault')}/{qn('w:rPr')}")
    if rpr_default is not None:
        rf = rpr_default.find(qn("w:rFonts"))
        if rf is None:
            rf = rpr_default.makeelement(qn("w:rFonts"), {})
            rpr_default.insert(0, rf)
        for attr in list(rf.attrib):
            del rf.attrib[attr]
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(attr), a.font)
    for name, size, color in (("Author", a.size + 1, None), ("Date", a.size, None), ("Abstract", a.size - 1, None),
                              ("Footnote Text", a.size - 2, None), ("Block Text", a.size, None),
                              ("TOC Heading", round(a.size * 1.45), a.heading_color)):
        if name in by_name:
            set_font(by_name[name], hfont if name == "TOC Heading" else a.font, size, color,
                     bold=True if name == "TOC Heading" else None)

    for s in d.sections:
        if a.a4:
            s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(a.margins_cm)

    d.save(a.output)
    print(f"Wrote {a.output}. Use: python docx_convert.py in.md --to docx --reference-doc {a.output}", file=sys.stderr)


if __name__ == "__main__":
    main()
