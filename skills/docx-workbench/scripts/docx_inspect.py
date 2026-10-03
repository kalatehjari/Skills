#!/usr/bin/env python3
"""Summarise a .docx before working on it.

Reports: core properties, word/paragraph counts, heading outline, styles in use,
tables (size + first row), images, sections (page size, orientation, margins),
headers/footers, comments, tracked changes, fields (TOC, page numbers, SEQ),
and {{placeholders}} found anywhere in the document.

Usage:
    python docx_inspect.py report.docx [--json] [--outline-depth 3]
"""
import argparse
import json
import os
import re
import sys
from collections import Counter

sys.dont_write_bytecode = True
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from _docx_common import iter_paragraphs  # noqa: E402

import docx  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

EMU_PER_CM = 360000
PLACEHOLDER = re.compile(r"\{\{\s*([^{}]+?)\s*\}\}")


def heading_level(p):
    name = (p.style.name if p.style is not None else "") or ""
    if name == "Title":
        return 0
    m = re.match(r"Heading (\d)", name)
    if m:
        return int(m.group(1))
    lvl = p._p.find(f"{qn('w:pPr')}/{qn('w:outlineLvl')}")
    if lvl is not None:
        return int(lvl.get(qn("w:val"))) + 1
    return None


def inspect(path, depth):
    d = docx.Document(path)
    body = d.element.body
    info = {"file": path}
    cp = d.core_properties
    info["properties"] = {k: str(getattr(cp, k)) for k in ("title", "author", "subject", "keywords", "last_modified_by", "created", "modified", "revision") if getattr(cp, k)}

    paras = list(d.paragraphs)
    words = sum(len(p.text.split()) for p in iter_paragraphs(d, headers=False))
    info["counts"] = {"paragraphs_body": len(paras), "words_incl_tables": words, "tables": len(d.tables),
                      "inline_images": len(d.inline_shapes), "sections": len(d.sections)}

    outline = []
    for p in paras:
        lvl = heading_level(p)
        if lvl is not None and lvl <= depth and p.text.strip():
            outline.append({"level": lvl, "text": p.text.strip()[:120]})
    info["outline"] = outline
    info["styles_used"] = dict(Counter(p.style.name for p in iter_paragraphs(d) if p.style is not None).most_common(15))

    info["tables_detail"] = []
    for i, t in enumerate(d.tables, start=1):
        first = [c.text.strip()[:25] for c in t.rows[0].cells] if t.rows else []
        info["tables_detail"].append({"index": i, "rows": len(t.rows), "cols": len(t.columns), "style": t.style.name if t.style else None, "first_row": first})

    secs = []
    for i, s in enumerate(d.sections, start=1):
        secs.append({
            "section": i,
            "page_cm": [round(s.page_width / EMU_PER_CM, 1), round(s.page_height / EMU_PER_CM, 1)] if s.page_width else None,
            "orientation": "landscape" if s.orientation == 1 else "portrait",
            "margins_cm": [round(m / EMU_PER_CM, 2) if m is not None else None for m in (s.top_margin, s.right_margin, s.bottom_margin, s.left_margin)],
            "header": None if s.header.is_linked_to_previous else " | ".join(p.text for p in s.header.paragraphs if p.text)[:100],
            "footer": None if s.footer.is_linked_to_previous else " | ".join(p.text for p in s.footer.paragraphs if p.text)[:100],
            "different_first_page": bool(s.different_first_page_header_footer),
        })
    info["sections_detail"] = secs

    ins = body.findall(f".//{qn('w:ins')}")
    dels = body.findall(f".//{qn('w:del')}")
    authors = Counter(e.get(qn("w:author")) for e in ins + dels)
    info["tracked_changes"] = {"insertions": len(ins), "deletions": len(dels),
                               "format_changes": len(body.findall(f".//{qn('w:rPrChange')}")) + len(body.findall(f".//{qn('w:pPrChange')}")),
                               "authors": dict(authors)}
    try:
        comments = list(d.comments)
        info["comments"] = [{"author": c.author, "text": c.text[:120]} for c in comments][:50]
    except Exception:
        info["comments"] = [] if not body.findall(f".//{qn('w:commentReference')}") else ["(present; python-docx too old to read)"]

    instr = [e.text or "" for e in body.iter(qn("w:instrText"))] + [e.get(qn("w:instr")) or "" for e in body.iter(qn("w:fldSimple"))]
    kinds = Counter((s.strip().split() or ["?"])[0].upper() for s in instr if s.strip())
    info["fields"] = dict(kinds)
    M = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
    info["equations"] = len(body.findall(f".//{M}oMath"))
    info["footnotes"] = len(body.findall(f".//{qn('w:footnoteReference')}"))
    info["endnotes"] = len(body.findall(f".//{qn('w:endnoteReference')}"))
    info["text_boxes"] = len(body.findall(f".//{qn('w:txbxContent')}"))

    found = Counter()
    for p in iter_paragraphs(d):
        for m in PLACEHOLDER.finditer(p.text):
            found[m.group(1)] += 1
    info["placeholders"] = dict(found)
    return info


def print_human(i):
    print(f"File: {i['file']}")
    if i["properties"]:
        print("Properties: " + "; ".join(f"{k}={v}" for k, v in i["properties"].items()))
    c = i["counts"]
    print(f"Content: {c['words_incl_tables']} words, {c['paragraphs_body']} body paragraphs, {c['tables']} tables, {c['inline_images']} images, {c['sections']} section(s)")
    if i["outline"]:
        print("Outline:")
        for h in i["outline"]:
            print("  " + "  " * max(h["level"] - 1, 0) + ("[T] " if h["level"] == 0 else f"H{h['level']} ") + h["text"])
    else:
        print("Outline: no heading styles used (headings may be manual bold text)")
    print("Styles: " + ", ".join(f"{k}×{v}" for k, v in i["styles_used"].items()))
    for t in i["tables_detail"]:
        print(f"Table {t['index']}: {t['rows']}×{t['cols']} ({t['style']}) first row: {t['first_row']}")
    for s in i["sections_detail"]:
        size = s["page_cm"] or [0, 0]
        paper = "A4" if sorted(size) == [21.0, 29.7] else "US Letter" if sorted(size) == [21.6, 27.9] else "custom"
        print(f"Section {s['section']}: {paper} {s['page_cm']} cm {s['orientation']}, margins T/R/B/L {s['margins_cm']} cm"
              + (f"; header '{s['header']}'" if s["header"] else "") + (f"; footer '{s['footer']}'" if s["footer"] else ""))
    tc = i["tracked_changes"]
    if tc["insertions"] or tc["deletions"] or tc["format_changes"]:
        print(f"Tracked changes: {tc['insertions']} insertions, {tc['deletions']} deletions, {tc['format_changes']} formatting; authors {tc['authors']}")
    else:
        print("Tracked changes: none")
    print(f"Comments: {len(i['comments'])}")
    if i["fields"]:
        print("Fields: " + ", ".join(f"{k}×{v}" for k, v in i["fields"].items()))
    extras = [f"{i[k]} {k.replace('_', ' ')}" for k in ("equations", "footnotes", "endnotes", "text_boxes") if i.get(k)]
    if extras:
        print("Also contains: " + ", ".join(extras) + (" (text-box text is not covered by replace/fill)" if i.get("text_boxes") else ""))
    if i["placeholders"]:
        print("Placeholders: " + ", ".join(f"{{{{{k}}}}}" + (f"×{v}" if v > 1 else "") for k, v in i["placeholders"].items()))


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("docx")
    ap.add_argument("--json", action="store_true")
    ap.add_argument("--outline-depth", type=int, default=3)
    a = ap.parse_args()
    if a.docx.lower().endswith(".doc"):
        sys.exit("Legacy .doc: convert first — python docx_convert.py file.doc --to docx")
    info = inspect(a.docx, a.outline_depth)
    print(json.dumps(info, indent=2, ensure_ascii=False, default=str) if a.json else "", end="")
    if not a.json:
        print_human(info)


if __name__ == "__main__":
    main()
