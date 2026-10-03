"""Shared helpers for the docx-workbench scripts (imported from the same folder).

- iter_paragraphs(doc): every paragraph in body, tables (incl. nested), headers, footers
- replace_in_paragraph(p, regex, repl): run-aware replace that survives Word splitting
  text such as "{{client_name}}" across several runs
"""
import re
import sys

sys.dont_write_bytecode = True  # keep the skill folder free of __pycache__

try:
    import docx  # noqa: F401
    from docx.text.paragraph import Paragraph
    from docx.text.run import Run
    from docx.oxml.ns import qn
except ImportError:
    sys.exit("python-docx is required: pip install python-docx")


def _walk(container, seen):
    for p in container.paragraphs:
        yield p
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                tc = cell._tc
                if tc in seen:  # merged cells repeat in row.cells; keep the element itself so identity is stable
                    continue
                seen.add(tc)
                yield from _walk(cell, seen)


def iter_paragraphs(doc, headers=True):
    seen = set()
    yield from _walk(doc, seen)
    if not headers:
        return
    parts = set()
    for s in doc.sections:
        for hf in (s.header, s.first_page_header, s.even_page_header,
                   s.footer, s.first_page_footer, s.even_page_footer):
            if hf.is_linked_to_previous:
                continue
            key = id(hf.part)
            if key in parts:
                continue
            parts.add(key)
            yield from _walk(hf, seen)


def text_runs(p):
    """Runs that carry the paragraph's visible text, in order (includes runs inside hyperlinks)."""
    runs = []
    for child in p._p.iterchildren():
        if child.tag == qn("w:r"):
            runs.append(Run(child, p))
        elif child.tag in (qn("w:hyperlink"), qn("w:ins"), qn("w:smartTag"), qn("w:fldSimple")):
            for r in child.iterchildren(qn("w:r")):
                runs.append(Run(r, p))
    return runs


def replace_in_paragraph(p, regex, repl):
    """Replace regex matches across run boundaries. `repl` is a str or a function(match)->str.

    The replacement inherits the formatting of the run where the match starts.
    Returns the number of replacements made.
    """
    runs = text_runs(p)
    if not runs:
        return 0
    texts = [r.text for r in runs]
    full = "".join(texts)
    matches = list(regex.finditer(full))
    if not matches:
        return 0
    # starting offset of each run
    starts, pos = [], 0
    for t in texts:
        starts.append(pos)
        pos += len(t)

    def locate(offset, end=False):
        for i in range(len(runs)):
            lo, hi = starts[i], starts[i] + len(texts[i])
            if (lo <= offset < hi) or (end and lo < offset <= hi):
                return i, offset - lo
        return len(runs) - 1, len(texts[-1])

    for m in reversed(matches):
        s, e = m.span()
        if s == e:
            continue
        new = repl(m) if callable(repl) else m.expand(repl)
        si, so = locate(s)
        ei, eo = locate(e, end=True)
        if si == ei:
            t = texts[si]
            texts[si] = t[:so] + new + t[eo:]
        else:
            texts[si] = texts[si][:so] + new
            for k in range(si + 1, ei):
                texts[k] = ""
            texts[ei] = texts[ei][eo:]
    for r, t in zip(runs, texts):
        if r.text != t:
            r.text = t
    return len([m for m in matches if m.start() != m.end()])


def compile_pattern(find, regex=False, ignore_case=False, whole_word=False):
    pat = find if regex else re.escape(find)
    if whole_word:
        pat = r"\b" + pat + r"\b"
    return re.compile(pat, re.IGNORECASE if ignore_case else 0)


# Order of children in word/settings.xml (CT_Settings). Word rejects out-of-order elements,
# so new settings must be inserted before the first element that the schema puts after them.
SETTINGS_ORDER = """writeProtection view zoom removePersonalInformation removeDateAndTime
doNotDisplayPageBoundaries displayBackgroundShape printPostScriptOverText printFractionalCharacterWidth
printFormsData embedTrueTypeFonts embedSystemFonts saveSubsetFonts saveFormsData mirrorMargins
alignBordersAndEdges bordersDoNotSurroundHeader bordersDoNotSurroundFooter gutterAtTop hideSpellingErrors
hideGrammaticalErrors activeWritingStyle proofState formsDesign attachedTemplate linkStyles
stylePaneFormatFilter stylePaneSortMethod documentType mailMerge revisionView trackRevisions
doNotTrackMoves doNotTrackFormatting documentProtection autoFormatOverride styleLockTheme styleLockQFSet
defaultTabStop autoHyphenation consecutiveHyphenLimit hyphenationZone doNotHyphenateCaps showEnvelope
summaryLength clickAndTypeStyle defaultTableStyle evenAndOddHeaders bookFoldRevPrinting bookFoldPrinting
bookFoldPrintingSheets drawingGridHorizontalSpacing drawingGridVerticalSpacing
displayHorizontalDrawingGridEvery displayVerticalDrawingGridEvery doNotUseMarginsForDrawingGridOrigin
drawingGridHorizontalOrigin drawingGridVerticalOrigin doNotShadeFormData noPunctuationKerning
characterSpacingControl printTwoOnOne strictFirstAndLastChars noLineBreaksAfter noLineBreaksBefore
savePreviewPicture doNotValidateAgainstSchema saveInvalidXml ignoreMixedContent alwaysShowPlaceholderText
doNotDemarcateInvalidXml saveXmlDataOnly useXSLTWhenSaving saveThroughXslt showXMLTags
alwaysMergeEmptyNamespace updateFields hdrShapeDefaults footnotePr endnotePr compat docVars rsids mathPr
attachedSchema themeFontLang clrSchemeMapping doNotIncludeSubdocsInStats doNotAutoCompressPictures
forceUpgrade captions readModeInkLockDown smartTagType schemaLibrary shapeDefaults doNotEmbedSmartTags
decimalSymbol listSeparator""".split()


def set_setting(doc, name, val=None):
    """Add (or update) a w:<name> element in settings.xml at its schema position."""
    from docx.oxml import OxmlElement
    settings = doc.settings.element
    el = settings.find(qn(f"w:{name}"))
    if el is None:
        el = OxmlElement(f"w:{name}")
        later = set(SETTINGS_ORDER[SETTINGS_ORDER.index(name) + 1:])
        anchor = next((c for c in settings if c.tag.split("}")[-1] in later), None)
        if anchor is not None:
            anchor.addprevious(el)
        else:
            settings.append(el)
    if val is not None:
        el.set(qn("w:val"), val)
    return el


def add_field(paragraph, instr, placeholder="1"):
    """Append a complex field (PAGE, NUMPAGES, SEQ Figure, TOC, PAGEREF ...) to a paragraph.
    `placeholder` is the cached result shown until Word updates fields."""
    from docx.oxml import OxmlElement

    def fld(kind):
        r = paragraph.add_run()
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
        r._r.append(el)
        return r
    fld("begin")
    r = paragraph.add_run()
    it = OxmlElement("w:instrText")
    it.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    it.text = f" {instr} "
    r._r.append(it)
    fld("separate")
    res = paragraph.add_run(placeholder)
    fld("end")
    return res


def ensure_paragraph_style(doc, style_id, name, based_on="Normal", font_size=None, bold=None,
                           color=None, space_before=None, space_after=None, tab_right_twips=None, indent_twips=None):
    """Create a paragraph style with an explicit Word styleId (e.g. 'TOC1', 'Footer') if missing."""
    from docx.oxml import OxmlElement
    styles_el = doc.styles.element
    for st in styles_el.iterchildren(qn("w:style")):
        if st.get(qn("w:styleId")) == style_id:
            return st
    st = OxmlElement("w:style")
    st.set(qn("w:type"), "paragraph")
    st.set(qn("w:styleId"), style_id)
    for tag, val in (("w:name", name), ("w:basedOn", based_on), ("w:uiPriority", "39")):
        e = OxmlElement(tag)
        e.set(qn("w:val"), val)
        st.append(e)
    st.append(OxmlElement("w:unhideWhenUsed"))
    ppr = OxmlElement("w:pPr")
    if tab_right_twips:
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "right")
        tab.set(qn("w:leader"), "dot")
        tab.set(qn("w:pos"), str(int(tab_right_twips)))
        tabs.append(tab)
        ppr.append(tabs)
    if space_before is not None or space_after is not None:
        sp = OxmlElement("w:spacing")
        if space_before is not None:
            sp.set(qn("w:before"), str(int(space_before * 20)))
        if space_after is not None:
            sp.set(qn("w:after"), str(int(space_after * 20)))
        ppr.append(sp)
    if indent_twips:
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), str(int(indent_twips)))
        ppr.append(ind)
    st.append(ppr)
    rpr = OxmlElement("w:rPr")
    if bold:
        rpr.append(OxmlElement("w:b"))
    if color:
        c = OxmlElement("w:color")
        c.set(qn("w:val"), color)
        rpr.append(c)
    if font_size:
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(int(font_size * 2)))
        rpr.append(sz)
    st.append(rpr)
    styles_el.append(st)
    return st


def set_style_id(paragraph, style_id):
    paragraph._p.get_or_add_pPr().get_or_add_pStyle().val = style_id
