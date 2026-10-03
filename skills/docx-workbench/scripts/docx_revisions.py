#!/usr/bin/env python3
"""List, accept or reject tracked changes; list comments with the text they refer to.

Subcommands:
  list      in.docx [--json]          every insertion/deletion: type, author, date, text
  comments  in.docx [--json]          every comment: author, date, comment text, anchored text
  accept    in.docx -o out.docx [--author NAME]   accept all (or only NAME's) changes
  reject    in.docx -o out.docx [--author NAME]   reject all (or only NAME's) changes

Handles run insertions/deletions, moves, inserted/deleted paragraph marks,
inserted/deleted table rows and formatting changes, in the body, headers,
footers, footnotes and endnotes. Comments are kept unless you add --drop-comments,
which deletes them completely (anchors, comment text and the comments part).
Options for accept/reject: -o OUT  --author NAME  --drop-comments  --stop-tracking
"""
import argparse
import json
import os
import sys

sys.dont_write_bytecode = True
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import _docx_common  # noqa: E402,F401

import docx  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from lxml import etree  # noqa: E402

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
STORY_TYPES = ("document.main", "header", "footer", "footnotes", "endnotes")


def story_roots(doc):
    """(part, root_element, is_blob) for every story part that can hold revisions."""
    out = []
    for part in doc.part.package.iter_parts():
        ct = part.content_type or ""
        if "wordprocessingml" not in ct or not any(s in ct for s in STORY_TYPES):
            continue
        if hasattr(part, "_element"):
            out.append((part, part._element, False))
        else:
            out.append((part, etree.fromstring(part.blob), True))
    return out


def text_of(el):
    return "".join(t.text or "" for t in el.iter(qn("w:t"), qn("w:delText"), qn("w:instrText")))


def unwrap(el):
    parent = el.getparent()
    idx = parent.index(el)
    for child in list(el):
        parent.insert(idx, child)
        idx += 1
    parent.remove(el)


def remove(el):
    el.getparent().remove(el)


def merge_with_next(p):
    """Remove p's paragraph mark: append the next paragraph's content to p, delete the next one."""
    nxt = p.getnext()
    while nxt is not None and nxt.tag != qn("w:p"):
        nxt = nxt.getnext()
    if nxt is None:
        return
    for child in list(nxt):
        if child.tag != qn("w:pPr"):
            p.append(child)
    remove(nxt)


def by(author, el):
    return author is None or el.get(qn("w:author")) == author


def resolve(root, mode, author):
    """mode = 'accept' or 'reject'. Returns number of revisions resolved."""
    n = 0
    keep_tag, drop_tag = ("w:ins", "w:del") if mode == "accept" else ("w:del", "w:ins")
    keep_move, drop_move = ("w:moveTo", "w:moveFrom") if mode == "accept" else ("w:moveFrom", "w:moveTo")

    # table rows inserted/deleted (marker lives in w:trPr)
    for mk in list(root.iter(qn("w:ins"), qn("w:del"))):
        par = mk.getparent()
        if par is not None and par.tag == qn("w:trPr") and by(author, mk):
            row = par.getparent()
            if mk.tag == qn(drop_tag):
                remove(row)
            else:
                remove(mk)
            n += 1

    # paragraph marks inserted/deleted (marker lives in w:pPr/w:rPr)
    for mk in list(root.iter(qn("w:ins"), qn("w:del"))):
        par = mk.getparent()
        if par is None or par.tag != qn("w:rPr") or par.getparent() is None or par.getparent().tag != qn("w:pPr"):
            continue
        if not by(author, mk):
            continue
        p = par.getparent().getparent()
        remove(mk)
        if mk.tag == qn(drop_tag):
            merge_with_next(p)
        n += 1

    # run-level content
    for el in list(root.iter(qn(drop_tag), qn(drop_move))):
        if el.getparent() is not None and by(author, el) and el.getparent().tag not in (qn("w:rPr"), qn("w:trPr")):
            remove(el)
            n += 1
    for el in list(root.iter(qn(keep_tag), qn(keep_move))):
        if el.getparent() is not None and by(author, el) and el.getparent().tag not in (qn("w:rPr"), qn("w:trPr")):
            for t in el.iter(qn("w:delText")):
                t.tag = qn("w:t")
            for t in el.iter(qn("w:delInstrText")):
                t.tag = qn("w:instrText")
            unwrap(el)
            n += 1
    for tag in ("w:moveFromRangeStart", "w:moveFromRangeEnd", "w:moveToRangeStart", "w:moveToRangeEnd"):
        for el in list(root.iter(qn(tag))):
            remove(el)

    # formatting changes
    for tag in ("w:rPrChange", "w:pPrChange", "w:sectPrChange", "w:tblPrChange", "w:trPrChange", "w:tcPrChange", "w:tblGridChange", "w:numberingChange"):
        for ch in list(root.iter(qn(tag))):
            if not by(author, ch):
                continue
            if mode == "reject":
                owner = ch.getparent()  # e.g. w:rPr; the change holds the OLD properties
                old = ch.find(f"{{{W}}}{owner.tag.split('}')[1]}")
                for child in list(owner):
                    if child is not ch:
                        owner.remove(child)
                if old is not None:
                    for child in list(old):
                        owner.append(child)
            remove(ch)
            n += 1
    return n


def para_text(p, accepted=True):
    """Paragraph text as it would read with revisions accepted (or rejected)."""
    out = []
    for node in p.iter(qn("w:t"), qn("w:delText"), qn("w:tab")):
        anc = node.getparent()
        in_del = in_ins = False
        while anc is not None and anc is not p:
            if anc.tag in (qn("w:del"), qn("w:moveFrom")):
                in_del = True
            if anc.tag in (qn("w:ins"), qn("w:moveTo")):
                in_ins = True
            anc = anc.getparent()
        if (accepted and in_del) or (not accepted and in_ins):
            continue
        out.append("\t" if node.tag == qn("w:tab") else (node.text or ""))
    return "".join(out)


def style_of(p):
    ps = p.find(f"{qn('w:pPr')}/{qn('w:pStyle')}")
    return ps.get(qn("w:val")) if ps is not None else ""


def cmd_list(a):
    doc = docx.Document(a.docx)
    rows = []
    for part, root, _ in story_roots(doc):
        where = os.path.basename(str(part.partname))
        heading, pno = "", 0
        for p in root.iter(qn("w:p")):
            pno += 1
            st = style_of(p).lower()
            if st.startswith("heading") or st == "title":
                heading = para_text(p)[:60]
            context = para_text(p)
            changes = []
            for el in p.iter(qn("w:ins"), qn("w:del"), qn("w:moveFrom"), qn("w:moveTo"), qn("w:rPrChange"), qn("w:pPrChange")):
                par = el.getparent()
                kind = el.tag.split("}")[1]
                if kind in ("rPrChange", "pPrChange"):
                    kind = "format"
                elif par is not None and par.tag == qn("w:rPr") and par.getparent() is not None and par.getparent().tag == qn("w:pPr"):
                    kind += "-paragraph-mark"
                elif par is not None and par.tag in (qn("w:rPr"), qn("w:trPr")):
                    continue
                changes.append({"el": el, "type": kind, "author": el.get(qn("w:author")), "date": (el.get(qn("w:date")) or "")[:10],
                                "text": text_of(el) if kind != "format" else ""})
            i = 0
            while i < len(changes):
                c = changes[i]
                nxt = changes[i + 1] if i + 1 < len(changes) else None
                row = {"part": where, "paragraph": pno, "heading": heading, "author": c["author"], "date": c["date"]}
                if (c["type"] == "del" and nxt and nxt["type"] == "ins" and nxt["author"] == c["author"]
                        and c["el"].getnext() is nxt["el"]):
                    row.update(type="replace", old=c["text"], new=nxt["text"])
                    i += 2
                else:
                    row.update(type=c["type"], text=c["text"])
                    i += 1
                row["context"] = context[:160]
                rows.append(row)
    if a.json:
        print(json.dumps(rows, indent=2, ensure_ascii=False))
        return
    if not rows:
        print("No tracked changes.")
    for r in rows:
        what = f"{r['old']!r} → {r['new']!r}" if r["type"] == "replace" else repr(r.get("text", ""))
        loc = f"¶{r['paragraph']}" + (f" under '{r['heading']}'" if r["heading"] else "") + (f" [{r['part']}]" if r["part"] != "document.xml" else "")
        print(f"[{r['type']}] {what}  — {r['author']} {r['date']}, {loc}\n    reads (accepted): {r['context']!r}")
    print(f"{len(rows)} change(s)", file=sys.stderr)


def cmd_comments(a):
    doc = docx.Document(a.docx)
    anchors = {}
    body = doc.element.body
    active = {}  # comment id -> [visible text, deleted text]
    for node in body.iter():
        if node.tag == qn("w:commentRangeStart"):
            active[node.get(qn("w:id"))] = [[], []]
        elif node.tag == qn("w:commentRangeEnd"):
            cid = node.get(qn("w:id"))
            if cid in active:
                vis, dele = active.pop(cid)
                vis_t, del_t = "".join(vis), "".join(dele)
                anchors[cid] = vis_t + (f" [deleted: {del_t}]" if del_t else "")
        elif node.tag == qn("w:t"):
            for v in active.values():
                v[0].append(node.text or "")
        elif node.tag == qn("w:delText"):
            for v in active.values():
                v[1].append(node.text or "")
    out = []
    try:
        for c in doc.comments:
            out.append({"id": c.comment_id, "author": c.author, "initials": c.initials,
                        "date": str(getattr(c, "timestamp", "") or ""), "comment": c.text,
                        "anchored_text": anchors.get(str(c.comment_id), "")})
    except AttributeError:
        sys.exit("Reading comments needs python-docx >= 1.2: pip install -U python-docx")
    if a.json:
        print(json.dumps(out, indent=2, ensure_ascii=False))
        return
    if not out:
        print("No comments.")
    for c in out:
        on = c["anchored_text"][:150] or "(anchor text no longer present — e.g. the commented text was deleted)"
        print(f"#{c['id']} {c['author']}: {c['comment']}\n    on: {on!r}")


COMMENT_RELS = ("/comments", "/commentsExtended", "/commentsIds", "/commentsExtensible", "/people")


def drop_comments(doc):
    """Remove comment anchors from the body AND the comment parts themselves,
    so no comment text remains anywhere in the saved file."""
    for tag in ("w:commentRangeStart", "w:commentRangeEnd"):
        for el in list(doc.element.body.iter(qn(tag))):
            remove(el)
    for r in list(doc.element.body.iter(qn("w:r"))):
        if r.find(qn("w:commentReference")) is not None:
            remove(r)
    rels = doc.part.rels
    for rid in [rid for rid, rel in rels.items() if rel.reltype.endswith(COMMENT_RELS)]:
        rels.pop(rid)  # parts no longer referenced are not written on save


def cmd_resolve(a, mode):
    if os.path.abspath(a.output) == os.path.abspath(a.docx):
        sys.exit("Refusing to overwrite the input; choose a different -o")
    doc = docx.Document(a.docx)
    total = 0
    for part, root, is_blob in story_roots(doc):
        n = resolve(root, mode, a.author)
        total += n
        if is_blob and n:
            part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    if a.drop_comments:
        drop_comments(doc)
    settings = doc.settings.element
    tr = settings.find(qn("w:trackRevisions"))
    if tr is not None and a.stop_tracking:
        settings.remove(tr)
    doc.save(a.output)
    print(f"Wrote {a.output}: {mode}ed {total} revision(s)" + (f" by {a.author}" if a.author else ""), file=sys.stderr)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    for name in ("list", "comments"):
        s = sub.add_parser(name)
        s.add_argument("docx")
        s.add_argument("--json", action="store_true")
    for name in ("accept", "reject"):
        s = sub.add_parser(name)
        s.add_argument("docx")
        s.add_argument("-o", "--output", required=True)
        s.add_argument("--author", help="only resolve changes by this author")
        s.add_argument("--drop-comments", action="store_true", help="also delete all comments (anchors and comment text) — use for a clean copy to send out")
        s.add_argument("--stop-tracking", action="store_true", help="switch off Track Changes in the output")
    a = ap.parse_args()
    if a.cmd == "list":
        cmd_list(a)
    elif a.cmd == "comments":
        cmd_comments(a)
    else:
        cmd_resolve(a, a.cmd)


if __name__ == "__main__":
    main()
