#!/usr/bin/env python3
"""Extract a deck's content as Markdown: slide titles, bullet hierarchy, tables,
chart data, picture alt text and speaker notes — for summarising, reviewing or
rebuilding a deck.

Usage:
    python pptx_text.py deck.pptx [-o deck.md] [--slides 2-6] [--no-notes]
"""
import argparse
import os
import sys

sys.dont_write_bytecode = True
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from _pptx_common import iter_shapes, parse_slides, slide_title  # noqa: E402

from pptx import Presentation  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE_TYPE  # noqa: E402


def is_chrome(sh):
    """Footers, slide numbers and dates: repeated on every slide, not content."""
    from pptx.enum.shapes import PP_PLACEHOLDER
    if sh.is_placeholder and sh.placeholder_format.type in (PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.DATE):
        return True
    return sh.name in ("Footer", "Slide Number")


def shape_md(sh, title_text):
    out = []
    if getattr(sh, "has_table", False) and sh.has_table:
        rows = [[c.text.replace("\n", " ").replace("|", "\\|") for c in r.cells] for r in sh.table.rows]
        if rows:
            out.append("| " + " | ".join(rows[0]) + " |")
            out.append("|" + "---|" * len(rows[0]))
            out += ["| " + " | ".join(r) + " |" for r in rows[1:]]
        return out
    if getattr(sh, "has_chart", False) and sh.has_chart:
        ch = sh.chart
        out.append(f"*[Chart: {ch.chart_type}".split(" (")[0] + "]*")
        try:
            cats = list(ch.plots[0].categories)
            out.append("| Series | " + " | ".join(str(c) for c in cats) + " |")
            out.append("|" + "---|" * (len(cats) + 1))
            for s in ch.series:
                out.append(f"| {s.name} | " + " | ".join(f"{v:g}" if isinstance(v, (int, float)) else str(v) for v in s.values) + " |")
        except Exception:
            pass
        return out
    if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
        alt = sh._element.nvPicPr.cNvPr.get("descr") or sh.name
        out.append(f"*[Picture: {alt}]*")
        return out
    if sh.has_text_frame:
        counters = {}
        for p in sh.text_frame.paragraphs:
            t = "".join(r.text for r in p.runs).strip()
            if not t or t == title_text:
                continue
            pPr = p._p.pPr
            numbered = pPr is not None and pPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}buAutoNum") is not None
            plain = pPr is not None and pPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}buNone") is not None
            if numbered:
                counters[p.level] = counters.get(p.level, 0) + 1
                for deeper in [k for k in counters if k > p.level]:
                    del counters[deeper]
            marker = f"{counters[p.level]}. " if numbered else ("" if plain or len(sh.text_frame.paragraphs) == 1 else "- ")
            out.append("  " * p.level + marker + t)
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("deck")
    ap.add_argument("-o", "--output")
    ap.add_argument("--slides")
    ap.add_argument("--no-notes", action="store_true")
    a = ap.parse_args()
    prs = Presentation(a.deck)
    lines = []
    for i in parse_slides(a.slides, len(prs.slides)):
        s = prs.slides[i]
        title = slide_title(s)
        lines.append(f"## Slide {i + 1}: {title}")
        lines.append("")
        shapes = sorted(iter_shapes(s.shapes), key=lambda x: (x.top or 0, x.left or 0))
        shapes = [x for x in shapes if not is_chrome(x)]
        for sh in shapes:
            md = shape_md(sh, title)
            if md:
                lines += md + [""]
        if not a.no_notes and s.has_notes_slide:
            notes = s.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines += ["> **Notes:** " + notes.replace("\n", " "), ""]
    text = "\n".join(lines).rstrip() + "\n"
    if a.output:
        with open(a.output, "w", encoding="utf-8") as fh:
            fh.write(text)
        print(f"Wrote {a.output}", file=sys.stderr)
    else:
        sys.stdout.write(text)


if __name__ == "__main__":
    main()
