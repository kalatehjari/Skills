#!/usr/bin/env python3
"""Quality check a deck before it is shown or sent.

Flags, per slide:
  overflow     text likely too long for its box (estimate from font size & box size)
  small-font   text below --min-font pt (default 14; 18+ recommended for body text in talks)
  dense        more than --max-words words (default 40, tables excluded) or --max-bullets (6) bullets
  off-slide    shapes partly outside the slide
  overlap      text boxes overlapping other text/pictures/tables
  empty        empty placeholders ("Click to add text" shows in edit view, may print)
  no-title     slide without a real title placeholder (outline view, screen readers and
               PowerPoint's accessibility checker use it; a text box that looks like a title is not enough)
  chart-axes   chart without axis titles
  duplicate    the same picture on several slides (often a placeholder figure never replaced)
  alt-text     pictures without alternative text
  low-res      pictures shown at < 120 px per inch
  leftover     {{placeholders}}, "Lorem ipsum", "TODO", "TBC", "XX" left in text
  fonts        more than 3 font families used

Usage:
    python pptx_check.py deck.pptx [--json] [--min-font 14] [--max-words 40] [--max-bullets 6]
Defaults follow references/design-principles.md (talk slides). For reading decks / handouts
use e.g. --max-words 120 --max-bullets 10.
"""
import argparse
import io
import json
import os
import re
import sys
from collections import Counter

sys.dont_write_bytecode = True
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from _pptx_common import iter_shapes, slide_title  # noqa: E402,F401

from pptx import Presentation  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER  # noqa: E402

LEFTOVER = re.compile(r"\{\{[^}]+\}\}|lorem ipsum|\bTODO\b|\bTBC\b|\bTBD\b|\bXX+\b|click to (add|edit)", re.I)


def font_size_pt(run, para, shape):
    for src in (run.font.size, getattr(para.font, "size", None)):
        if src:
            return src.pt
    return 18.0  # typical inherited body size when not set explicitly


def text_height_estimate(tf, width_pt, shape):
    h = 0.0
    for p in tf.paragraphs:
        text = "".join(r.text for r in p.runs)
        size = max([font_size_pt(r, p, shape) for r in p.runs] or [18.0])
        usable = max(width_pt - 14 - p.level * 20, 30)
        cpl = max(usable / (size * 0.5), 6)
        lines = max(1, -(-len(text) // int(cpl))) + text.count("\n")
        h += lines * size * 1.2 + size * 0.3
    return h


def boxes_overlap(a, b):
    ax0, ay0, ax1, ay1 = a.left, a.top, a.left + a.width, a.top + a.height
    bx0, by0, bx1, by1 = b.left, b.top, b.left + b.width, b.top + b.height
    ix = min(ax1, bx1) - max(ax0, bx0)
    iy = min(ay1, by1) - max(ay0, by0)
    if ix <= 0 or iy <= 0:
        return 0
    return ix * iy / max(1, min(a.width * a.height, b.width * b.height))


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("deck")
    ap.add_argument("--json", action="store_true")
    ap.add_argument("--min-font", type=float, default=14)
    ap.add_argument("--max-words", type=int, default=40)
    ap.add_argument("--max-bullets", type=int, default=6)
    a = ap.parse_args()
    prs = Presentation(a.deck)
    W, H = prs.slide_width, prs.slide_height
    findings = []
    families = Counter()
    pics = {}

    def add(i, check, msg):
        findings.append({"slide": i, "check": check, "note": msg})

    for i, s in enumerate(prs.slides, start=1):
        shapes = [sh for sh in iter_shapes(s.shapes) if sh.width is not None and sh.left is not None]
        words, bullets = 0, 0
        tph = s.shapes.title
        if tph is None:
            add(i, "no-title", "no title placeholder" + (f" (visible heading '{slide_title(s)[:40]}' is a plain text box)" if slide_title(s) else ""))
        elif not tph.text_frame.text.strip():
            add(i, "no-title", "title placeholder is empty")
        for sh in shapes:
            if sh.left < -12700 or sh.top < -12700 or sh.left + sh.width > W + 12700 or sh.top + sh.height > H + 12700:
                add(i, "off-slide", f"'{sh.name}' extends beyond the slide edge")
            if sh.is_placeholder and sh.has_text_frame and not sh.text_frame.text.strip() \
                    and sh.placeholder_format.type not in (PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER):
                add(i, "empty", f"empty placeholder '{sh.name}' — delete it or fill it")
            if sh.has_text_frame and sh.text_frame.text.strip():
                tf = sh.text_frame
                is_footer = sh.top > H * 0.88 and sh.height < H * 0.08
                if not is_footer:
                    words += len(tf.text.split())
                    if sh.name not in ("Title",) and not (sh.is_placeholder and sh.placeholder_format.type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)):
                        bullets += sum(1 for p in tf.paragraphs if "".join(r.text for r in p.runs).strip())
                need = text_height_estimate(tf, sh.width / 12700, sh)
                if need > sh.height / 12700 * 1.1 and not is_footer:
                    add(i, "overflow", f"'{sh.name}': text needs ≈{need:.0f} pt, box is {sh.height / 12700:.0f} pt high")
                for p in tf.paragraphs:
                    for r in p.runs:
                        if r.font.name:
                            families[r.font.name] += 1
                        if r.text.strip() and r.font.size and r.font.size.pt < a.min_font and not is_footer:
                            add(i, "small-font", f"'{r.text.strip()[:30]}' at {r.font.size.pt:g} pt")
                            break
                m = LEFTOVER.search(tf.text)
                if m:
                    add(i, "leftover", f"'{m.group(0)}' in '{sh.name}'")
            if getattr(sh, "has_chart", False) and sh.has_chart:
                ch = sh.chart
                try:
                    if ch.chart_type is not None and "PIE" not in str(ch.chart_type) and "DOUGHNUT" not in str(ch.chart_type):
                        missing = [n for n, ax in (("x", ch.category_axis), ("y", ch.value_axis)) if not ax.has_title]
                        if missing:
                            add(i, "chart-axes", f"chart '{sh.name}' has no {' or '.join(missing)} axis title (give units)")
                except Exception:
                    pass
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    import hashlib
                    pics.setdefault(hashlib.md5(sh.image.blob).hexdigest(), []).append(i)
                except Exception:
                    pass
            if getattr(sh, "has_table", False) and sh.has_table:
                if len(sh.table.rows) > 8:
                    add(i, "dense", f"table with {len(sh.table.rows)} rows — keep ≤ 8 on a slide, move the rest to a handout")
                for row in sh.table.rows:
                    for cell in row.cells:
                        m = LEFTOVER.search(cell.text)
                        if m:
                            add(i, "leftover", f"'{m.group(0)}' in a table cell")
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                descr = sh._element.nvPicPr.cNvPr.get("descr")
                if not descr:
                    add(i, "alt-text", f"picture '{sh.name}' has no alt text")
                try:
                    from PIL import Image
                    with Image.open(io.BytesIO(sh.image.blob)) as im:
                        px = im.size[0]
                    ppi = px / (sh.width / 914400)
                    if ppi < 120:
                        add(i, "low-res", f"picture '{sh.name}' shown at ≈{ppi:.0f} px/inch (blurry when projected/printed)")
                except Exception:
                    pass
        chrome = lambda sh: sh.name in ("Footer", "Slide Number") or (sh.is_placeholder and sh.placeholder_format.type in (PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.DATE))  # noqa: E731
        texty = [sh for sh in shapes if sh.has_text_frame and sh.text_frame.text.strip() and not chrome(sh)]
        solid = [sh for sh in shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE or getattr(sh, "has_table", False) and sh.has_table
                 or getattr(sh, "has_chart", False) and sh.has_chart or sh.has_text_frame and sh.text_frame.text.strip()]
        for t in texty:
            for o in solid:
                if o is t or o.shape_id == t.shape_id:
                    continue
                if boxes_overlap(t, o) > 0.15:
                    add(i, "overlap", f"'{t.name}' overlaps '{o.name}'")
        if words > a.max_words:
            add(i, "dense", f"{words} words (aim for ≤ {a.max_words}); split the slide or move detail to notes")
        if bullets > a.max_bullets:
            add(i, "dense", f"{bullets} bullet lines (aim for ≤ {a.max_bullets})")
    for digest, where in pics.items():
        if len(set(where)) > 1:
            findings.append({"slide": where[0], "check": "duplicate",
                             "note": f"same picture on slides {sorted(set(where))} — intended, or a placeholder figure?"})
    if len(families) > 3:
        findings.append({"slide": 0, "check": "fonts", "note": f"{len(families)} font families: {dict(families)}"})
    # de-duplicate overlap pairs
    seen, uniq = set(), []
    for f in findings:
        key = (f["slide"], f["check"], frozenset(re.findall(r"'([^']+)'", f["note"])) if f["check"] == "overlap" else f["note"])
        if key not in seen:
            seen.add(key)
            uniq.append(f)
    findings = uniq
    summary = Counter(f["check"] for f in findings)
    if a.json:
        print(json.dumps({"summary": summary, "findings": findings}, indent=2, ensure_ascii=False))
        return
    print(f"Checked {len(prs.slides)} slides: " + (", ".join(f"{v} {k}" for k, v in summary.items()) or "no issues found"))
    for f in sorted(findings, key=lambda f: (f["slide"], f["check"])):
        print(f"  slide {f['slide']:>2}  {f['check']:<10} {f['note']}")
    print("Estimates are heuristic — confirm visually: python pptx_render.py DECK --sheet", file=sys.stderr)


if __name__ == "__main__":
    main()
