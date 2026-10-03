"""Shared helpers for pptx-workbench scripts (imported from the same folder)."""
import os
import re
import shutil
import subprocess
import sys
import tempfile

sys.dont_write_bytecode = True

try:
    from pptx import Presentation  # noqa: F401
    from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER  # noqa: F401
    from pptx.util import Emu, Pt  # noqa: F401
except ImportError:
    sys.exit("python-pptx is required: pip install python-pptx")

EMU_PER_PT = 12700
EMU_PER_IN = 914400


def iter_shapes(shapes):
    """All shapes, descending into groups."""
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(sh.shapes)
        else:
            yield sh


def iter_text_frames(slide, notes=False):
    """(shape, text_frame) for every text-bearing shape, table cell and (optionally) the notes."""
    for sh in iter_shapes(slide.shapes):
        if sh.has_text_frame:
            yield sh, sh.text_frame
        if getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                for cell in row.cells:
                    yield sh, cell.text_frame
    if notes and slide.has_notes_slide:
        yield None, slide.notes_slide.notes_text_frame


def replace_in_paragraph(p, regex, repl):
    """Run-aware replace inside one paragraph; keeps the formatting of the run where a match starts."""
    runs = list(p.runs)
    if not runs:
        return 0
    texts = [r.text for r in runs]
    full = "".join(texts)
    matches = [m for m in regex.finditer(full) if m.start() != m.end()]
    if not matches:
        return 0
    starts, pos = [], 0
    for t in texts:
        starts.append(pos)
        pos += len(t)

    def locate(off, end=False):
        for i in range(len(runs)):
            lo, hi = starts[i], starts[i] + len(texts[i])
            if (lo <= off < hi) or (end and lo < off <= hi):
                return i, off - lo
        return len(runs) - 1, len(texts[-1])

    for m in reversed(matches):
        new = repl(m) if callable(repl) else m.expand(repl)
        si, so = locate(m.start())
        ei, eo = locate(m.end(), end=True)
        if si == ei:
            texts[si] = texts[si][:so] + new + texts[si][eo:]
        else:
            texts[si] = texts[si][:so] + new
            for k in range(si + 1, ei):
                texts[k] = ""
            texts[ei] = texts[ei][eo:]
    for r, t in zip(runs, texts):
        if r.text != t:
            r.text = t
    if any("\n" in t or "\v" in t for t in texts):
        split_paragraph_lines(p)
    return len(matches)


def split_paragraph_lines(p):
    """Turn line breaks inside a paragraph's runs into separate paragraphs (e.g. one bullet each),
    copying the paragraph's properties and the first run's formatting."""
    import copy
    from pptx.oxml.ns import qn
    full = "".join(r.text for r in p.runs).replace("\v", "\n")
    lines = full.split("\n")
    p_el = p._p
    runs = list(p.runs)
    template_r = copy.deepcopy(runs[0]._r) if runs else None
    for r in runs[1:]:
        p_el.remove(r._r)
    if runs:
        runs[0].text = lines[0]
    anchor = p_el
    for line in lines[1:]:
        new_p = copy.deepcopy(p_el)
        for r in new_p.findall(qn("a:r")) + new_p.findall(qn("a:br")) + new_p.findall(qn("a:fld")):
            new_p.remove(r)
        if template_r is not None:
            r = copy.deepcopy(template_r)
            r.find(qn("a:t")).text = line
            end = new_p.find(qn("a:endParaRPr"))
            if end is not None:
                end.addprevious(r)
            else:
                new_p.append(r)
        anchor.addnext(new_p)
        anchor = new_p


def soffice():
    for name in ("soffice", "libreoffice", "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        found = shutil.which(name) or (name if os.path.exists(name) else None)
        if found:
            return found
    return None


def to_pdf(src, outdir):
    exe = soffice()
    if not exe:
        sys.exit("LibreOffice not found (apt install libreoffice-impress / brew install --cask libreoffice)")
    with tempfile.TemporaryDirectory() as profile:
        res = subprocess.run([exe, f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "pdf",
                              "--outdir", outdir, src], capture_output=True, text=True, timeout=300)
    out = os.path.join(outdir, os.path.splitext(os.path.basename(src))[0] + ".pdf")
    if not os.path.exists(out):
        sys.exit(f"LibreOffice conversion failed: {res.stderr.strip() or res.stdout.strip()}")
    return out


def parse_slides(spec, n):
    """'1-3,7,10-' -> 0-based indices."""
    if not spec:
        return list(range(n))
    out = []
    for part in str(spec).split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            a, b = part.split("-", 1)
            s, e = (int(a) if a else 1), (int(b) if b else n)
        else:
            s = e = int(part)
        if s < 1 or e > n or s > e:
            sys.exit(f"Slide range '{part}' is outside 1-{n}")
        out += [i for i in range(s - 1, e) if i not in out]
    return out


def slide_title(slide):
    try:
        if slide.shapes.title is not None and slide.shapes.title.text.strip():
            return slide.shapes.title.text.strip()
    except Exception:
        pass
    for sh in iter_shapes(slide.shapes):
        if sh.has_text_frame and sh.text_frame.text.strip():
            return sh.text_frame.text.strip().split("\n")[0][:80]
    return ""


PLACEHOLDER = re.compile(r"\{\{\s*([A-Za-z0-9_.\-]+)\s*\}\}")
