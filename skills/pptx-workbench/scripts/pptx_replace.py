#!/usr/bin/env python3
"""Find/replace text across a deck, or fill {{placeholders}} from JSON, keeping formatting.

    python pptx_replace.py deck.pptx -o new.pptx --find "Lot 12" --replace "Lot 14"
    python pptx_replace.py deck.pptx -o new.pptx --map renames.json           # {"old": "new", ...}
    python pptx_replace.py template.pptx -o filled.pptx --fill data.json       # {{client}}, {{date}} ...
    python pptx_replace.py deck.pptx --find "draft" --dry-run                   # list matches per slide
Options: --regex  --ignore-case  --notes (also speaker notes)  --strict (fill: fail on missing keys)
Covers text boxes, placeholders, grouped shapes and table cells. Text inside charts is not changed.
A "\n" in a replacement (or a fill value) starts a new paragraph — e.g. one bullet per line:
    --find "Lorem ipsum dolor" --replace $'Findings\nOptions\nRecommendation'
(each new paragraph copies the original paragraph's bullet/level and the first run's formatting).
"""
import argparse
import json
import os
import re
import sys

sys.dont_write_bytecode = True
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from _pptx_common import PLACEHOLDER, iter_text_frames, replace_in_paragraph  # noqa: E402

from pptx import Presentation  # noqa: E402


def lookup(data, path):
    cur = data
    for part in path.split("."):
        if isinstance(cur, dict) and part in cur:
            cur = cur[part]
        else:
            return None
    return cur


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("deck")
    ap.add_argument("-o", "--output")
    ap.add_argument("--find")
    ap.add_argument("--replace", default="")
    ap.add_argument("--map")
    ap.add_argument("--fill")
    ap.add_argument("--regex", action="store_true")
    ap.add_argument("--ignore-case", action="store_true")
    ap.add_argument("--notes", action="store_true")
    ap.add_argument("--strict", action="store_true")
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()
    if not a.dry_run and not a.output:
        sys.exit("Give -o OUTPUT (or --dry-run)")
    if a.output and os.path.abspath(a.output) == os.path.abspath(a.deck):
        sys.exit("Refusing to overwrite the input; choose a different -o")
    prs = Presentation(a.deck)
    flags = re.I if a.ignore_case else 0

    jobs = []
    if a.map:
        with open(a.map, encoding="utf-8") as fh:
            jobs += [(re.compile(k if a.regex else re.escape(k), flags), v, k) for k, v in json.load(fh).items()]
    if a.find is not None:
        repl = a.replace if a.regex else a.replace.replace("\\", "\\\\")
        jobs.append((re.compile(a.find if a.regex else re.escape(a.find), flags), repl, a.find))
    data = None
    if a.fill:
        with open(a.fill, encoding="utf-8") as fh:
            data = json.load(fh)
    if not jobs and data is None:
        sys.exit("Give --find, --map or --fill")

    missing = set()
    totals = {}
    for i, slide in enumerate(prs.slides, start=1):
        for _, tf in iter_text_frames(slide, notes=a.notes):
            for p in tf.paragraphs:
                text = "".join(r.text for r in p.runs)
                for rx, repl, label in jobs:
                    if a.dry_run:
                        for m in rx.finditer(text):
                            print(f"  slide {i}: …{text[max(0, m.start() - 30):m.end() + 30]}…")
                            totals[label] = totals.get(label, 0) + 1
                    else:
                        totals[label] = totals.get(label, 0) + replace_in_paragraph(p, rx, repl)
                if data is not None and not a.dry_run:
                    def fill(m):
                        v = lookup(data, m.group(1))
                        if v is None:
                            missing.add(m.group(1))
                            return m.group(0)
                        return str(v)
                    totals["{{…}}"] = totals.get("{{…}}", 0) + replace_in_paragraph(p, PLACEHOLDER, fill)
    labels = [lab for _, _, lab in jobs] + (["{{…}}"] if data is not None else [])
    for label in labels:
        n = totals.get(label, 0)
        print(f"{label!r}: {n} {'match(es)' if a.dry_run else 'replacement(s)'}", file=sys.stderr)
    if missing:
        msg = f"No value for: {sorted(missing)}"
        if a.strict:
            sys.exit(msg)
        print("Warning: " + msg + " (left in place)", file=sys.stderr)
    if not a.dry_run:
        prs.save(a.output)
        print(f"Wrote {a.output}", file=sys.stderr)


if __name__ == "__main__":
    main()
