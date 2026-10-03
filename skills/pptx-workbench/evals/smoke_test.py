#!/usr/bin/env python3
"""Smoke test for pptx-workbench: builds, inspects, edits and renders sample decks.

    python evals/smoke_test.py
Render steps are skipped (with a note) when LibreOffice is missing.
"""
import json
import os
import shutil
import subprocess
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
SCRIPTS = os.path.join(os.path.dirname(HERE), "scripts")
F = lambda n: os.path.join(HERE, "files", n)  # noqa: E731
LO = shutil.which("soffice") or shutil.which("libreoffice")

OUTLINE = """---
title: Smoke test deck
author: Tester
footer: Smoke
---
# Section
## Bullets slide
- One
  - Sub
1. First
2. Second
Notes: say hello
## Figure
![Stress paths](stress_paths.png)
## Table
| A | B |
|---|---|
| 1 | 2.5 |
## Chart
```chart
{"type": "column", "categories": ["x", "y"], "series": {"s": [1, 2]}, "y_title": "v"}
```
## Two
::: left
- left text
:::
::: right
![Stress paths](stress_paths.png)
:::
## Statement
> Big message.
##
> Questions?
"""


def run(name, *args, expect=0):
    res = subprocess.run([sys.executable, os.path.join(SCRIPTS, name), *args], capture_output=True, text=True)
    ok = res.returncode == expect
    print(("PASS " if ok else "FAIL ") + name + " " + " ".join(os.path.basename(str(a)) for a in args))
    if not ok:
        print(res.stdout[-1500:], res.stderr[-1500:], sep="\n")
        sys.exit(1)
    return res


def main():
    with tempfile.TemporaryDirectory() as t:
        j = lambda *p: os.path.join(t, *p)  # noqa: E731
        shutil.copy(F("stress_paths.png"), j("stress_paths.png"))
        with open(j("talk.md"), "w") as fh:
            fh.write(OUTLINE)
        run("pptx_build.py", j("talk.md"), "-o", j("talk.pptx"))
        info = json.loads(run("pptx_inspect.py", j("talk.pptx"), "--json").stdout)
        assert info["slides"] == 9 and info["aspect"] == "16:9", info["slides"]
        md = run("pptx_text.py", j("talk.pptx")).stdout
        assert "say hello" in md and "| 1 | 2.5 |" in md and "Picture: Stress paths" in md
        chk = json.loads(run("pptx_check.py", j("talk.pptx"), "--json").stdout)["summary"]
        assert not chk.get("overflow") and not chk.get("alt-text") and not chk.get("no-title"), chk
        from pptx import Presentation
        assert all(sl.shapes.title is not None for sl in Presentation(j("talk.pptx")).slides), "title placeholders"
        run("pptx_replace.py", F("sample_template.pptx"), "-o", j("agenda.pptx"), "--find", "Lorem ipsum dolor",
            "--replace", "Findings\nOptions\nRecommendation")
        assert "- Findings\n- Options\n- Recommendation" in run("pptx_text.py", j("agenda.pptx")).stdout
        run("pptx_build.py", j("talk.md"), "-o", j("talk_t.pptx"), "--template", F("sample_template.pptx"))
        with open(j("fill.json"), "w") as fh:
            json.dump({"client": "ACME", "date": "1 Jan"}, fh)
        run("pptx_replace.py", F("sample_template.pptx"), "-o", j("filled.pptx"), "--fill", j("fill.json"), "--strict")
        assert "ACME" in run("pptx_text.py", j("filled.pptx")).stdout
        run("pptx_replace.py", F("sample_deck.pptx"), "-o", j("rep.pptx"), "--find", "Lot 12", "--replace", "Lot 14")
        run("pptx_slides.py", "reorder", j("talk.pptx"), "--order", "1,3,2,4-", "-o", j("ro.pptx"))
        run("pptx_slides.py", "delete", j("talk.pptx"), "--slides", "2", "-o", j("del.pptx"))
        run("pptx_slides.py", "duplicate", j("talk.pptx"), "--slides", "4", "-o", j("dup.pptx"))
        run("pptx_slides.py", "reorder", j("talk.pptx"), "--order", "1,2", "-o", j("bad.pptx"), expect=1)
        if LO:
            run("pptx_render.py", j("talk.pptx"), "--sheet", "-o", j("sheet.png"))
            run("pptx_render.py", j("talk.pptx"), "--pdf", "-o", j("talk.pdf"))
        else:
            print("(LibreOffice missing: skipped renders)")
    print("All checks passed.")


if __name__ == "__main__":
    main()
