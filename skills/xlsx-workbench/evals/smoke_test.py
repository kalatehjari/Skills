#!/usr/bin/env python3
"""Smoke test for xlsx-workbench: runs every script against the sample files.

    python evals/smoke_test.py
Steps needing LibreOffice are skipped (with a note) when it is missing.
"""
import json
import os
import shutil
import subprocess
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
SCRIPTS = os.path.join(os.path.dirname(HERE), "scripts")
F = lambda n: os.path.join(HERE, "files", n)  # noqa: E731
LO = shutil.which("soffice") or shutil.which("libreoffice")


def run(name, *args, expect=0):
    res = subprocess.run([sys.executable, os.path.join(SCRIPTS, name), *args], capture_output=True, text=True)
    ok = res.returncode == expect
    print(("PASS " if ok else "FAIL ") + name + " " + " ".join(os.path.basename(str(a)) for a in args))
    if not ok:
        print(res.stdout[-1500:], res.stderr[-1500:], sep="\n")
        sys.exit(1)
    return res


def main():
    with tempfile.TemporaryDirectory() as t:
        j = lambda *p: os.path.join(t, *p)  # noqa: E731
        info = json.loads(run("xlsx_inspect.py", F("sample_lab.xlsx"), "--json").stdout)
        assert info["sheets"][0]["header_row"] == 3
        csv_out = run("xlsx_read.py", F("sample_lab.xlsx"), "--format", "csv").stdout
        assert csv_out.startswith("Sample,Depth (m)")
        run("xlsx_from_data.py", F("sample_results.csv"), "-o", j("built.xlsx"), "--table", "--totals", "sum", "--total-cols", "Load (kN)")
        run("xlsx_add_chart.py", j("built.xlsx"), "-o", j("chart.xlsx"), "--type", "bar", "--x", "A2:A4", "--y", "E1:E4")
        run("xlsx_add_chart.py", F("sample_profile.xlsx"), "-o", j("prof.xlsx"), "--type", "scatter", "--x", "B2:B13,C2:C13",
            "--y", "A1:A13,A1:A13", "--series-names", "w (%),LL (%)", "--y-reverse", "--marker-only")
        run("xlsx_add_chart.py", F("sample_profile.xlsx"), "-o", j("bad.xlsx"), "--x", "B2:C13", "--y", "A1:A13", expect=1)
        run("xlsx_add_chart.py", F("sample_profile.xlsx"), "-o", F("sample_profile.xlsx"), "--x", "A2:A3", "--y", "B1:B3", expect=1)
        diff = run("xlsx_diff.py", F("sample_lab.xlsx"), j("built.xlsx")).stdout
        assert "Added sheets" in diff
        if LO:
            rec = json.loads(run("xlsx_recalc.py", F("sample_lab.xlsx"), "--json", "--cells", "Summary!B5").stdout)
            assert len(rec["errors"]) == 2 and rec["values"][0]["value"] == 5
            aud = json.loads(run("xlsx_audit.py", F("sample_lab.xlsx"), "--json").stdout)["summary"]
            assert aud.get("hardcoded") == 1 and aud.get("inconsistent") == 1 and aud.get("text-number") == 1
            assert aud.get("suspicious") == 1 and "magic-number" not in aud, aud
            tot = json.loads(run("xlsx_recalc.py", j("built.xlsx"), "--json", "--cells", "E5").stdout)["values"][0]["value"]
            assert abs(tot - 37320.75) < 1e-6, tot
            run("xlsx_recalc.py", j("chart.xlsx"), "--write", j("final.xlsx"))
            from openpyxl import load_workbook
            ws = load_workbook(j("final.xlsx"), data_only=True).worksheets[0]
            assert abs(ws["E5"].value - 37320.75) < 1e-6 and ws.tables, "cached value / table lost"
            run("xlsx_convert.py", F("sample_lab.xlsx"), "--to", "csv", "-o", j("csv"))
            run("xlsx_convert.py", j("chart.xlsx"), "--to", "png", "--fit", "-o", j("png"))
            run("xlsx_convert.py", j("chart.xlsx"), "--to", "ods", "-o", j("chart.ods"))
            run("xlsx_convert.py", j("chart.ods"), "--to", "xlsx", "-o", j("back.xlsx"))
        else:
            print("(LibreOffice missing: skipped recalc/audit/convert steps)")
    print("All checks passed.")


if __name__ == "__main__":
    main()
