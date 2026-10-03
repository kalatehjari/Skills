#!/usr/bin/env python3
"""Recalculate every formula with LibreOffice and report the results.

openpyxl writes formulas but cannot calculate them, and files saved by scripts
have no cached values; files from Excel may have stale ones. This script makes
a recalculated copy (LibreOffice, forced recalc-on-load), then reports:
  - formula errors (#DIV/0!, #REF!, #NAME?, #N/A, #VALUE! ...) with cell and formula
  - optionally the values of chosen cells (--cells "Summary!B5,Lab!F10")
The input file is not modified.
--write OUT   saves a copy of the ORIGINAL file with the computed results stored as
              cached values: every Excel feature is kept (table styles, charts,
              validation), and previews / pandas / data_only readers see numbers.
              This is the copy to hand to people.
--write-lo OUT saves LibreOffice's own re-saved workbook instead (may lose Excel-only
              features — see references/formulas-and-checking.md).

Usage:
    python xlsx_recalc.py model.xlsx [--cells "Sheet!A1,Sheet!B2:B5"] [--json]
                          [--write model_recalc.xlsx] [--fail-on-error]
"""
import argparse
import json
import os
import shutil
import sys
import tempfile

sys.dont_write_bytecode = True
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from _xlsx_common import ERROR_VALUES, inject_cached_values, recalculated_copy  # noqa: E402

from openpyxl import load_workbook  # noqa: E402


def cell_refs(spec, wb):
    for part in spec.split(","):
        part = part.strip()
        if not part:
            continue
        sheet, _, rng = part.rpartition("!")
        sheet = sheet.strip("'") or wb.sheetnames[0]
        ws = wb[sheet]
        cells = ws[rng]
        if not isinstance(cells, tuple):
            yield sheet, cells
        else:
            for row in cells if isinstance(cells[0], tuple) else (cells,):
                for c in row:
                    yield sheet, c


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx")
    ap.add_argument("--cells", action="append", help="cells/ranges to print, e.g. 'Summary!B2:C5,Lab!F10' (repeatable)")
    ap.add_argument("--json", action="store_true")
    ap.add_argument("--write", help="save the original with computed values cached (keeps all features)")
    ap.add_argument("--write-lo", help="save LibreOffice's re-saved workbook")
    ap.add_argument("--fail-on-error", action="store_true", help="exit 1 if any formula errors")
    a = ap.parse_args()
    for w in (a.write, a.write_lo):
        if w and os.path.abspath(w) == os.path.abspath(a.xlsx):
            sys.exit("Refusing to overwrite the input; choose a different output path")
    if a.write and os.path.splitext(a.xlsx)[1].lower() not in (".xlsx", ".xlsm"):
        sys.exit("--write needs an .xlsx/.xlsm input; use --write-lo for other formats")

    with tempfile.TemporaryDirectory() as tmp:
        rc = recalculated_copy(a.xlsx, tmp)
        wf = load_workbook(a.xlsx, data_only=False)
        wv = load_workbook(rc, data_only=True)
        errors, n_formulas = [], 0
        for ws in wf.worksheets:
            if ws.title not in wv.sheetnames:
                continue
            vs = wv[ws.title]
            for row in ws.iter_rows():
                for c in row:
                    if c.data_type == "f" or (isinstance(c.value, str) and str(c.value).startswith("=")):
                        n_formulas += 1
                        v = vs[c.coordinate].value
                        if isinstance(v, str) and v in ERROR_VALUES:
                            errors.append({"cell": f"{ws.title}!{c.coordinate}", "error": v, "formula": str(c.value)})
        values = []
        if a.cells:
            for sheet, c in cell_refs(",".join(a.cells), wv):
                f = wf[sheet][c.coordinate].value
                values.append({"cell": f"{sheet}!{c.coordinate}", "value": c.value,
                               "formula": str(f) if isinstance(f, str) and f.startswith("=") else None})
        if a.write:
            n_cached = inject_cached_values(a.xlsx, rc, a.write)
        if a.write_lo:
            shutil.copy(rc, a.write_lo)

    if a.json:
        print(json.dumps({"formulas": n_formulas, "errors": errors, "values": values}, indent=2, default=str, ensure_ascii=False))
    else:
        print(f"{n_formulas} formula(s) recalculated; {len(errors)} error(s)")
        for e in errors:
            print(f"  {e['cell']}: {e['error']}   {e['formula']}")
        for v in values:
            val = round(v["value"], 6) if isinstance(v["value"], float) else v["value"]
            print(f"  {v['cell']} = {val!r}" + (f"   ({v['formula']})" if v["formula"] else ""))
        if a.write:
            print(f"Wrote {a.write} ({n_cached} cached value(s) stored; original formatting and features kept)")
        if a.write_lo:
            print(f"Wrote {a.write_lo} (LibreOffice re-save)")
    if errors and a.fail_on_error:
        sys.exit(1)


if __name__ == "__main__":
    main()
