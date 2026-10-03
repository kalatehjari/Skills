#!/usr/bin/env python3
"""Cell-by-cell comparison of two workbooks (or two sheets).

Reports, per sheet: added/removed sheets, changed formulas, changed values
(constants, and formula results if --values), and rows/columns added at the end.

Usage:
    python xlsx_diff.py old.xlsx new.xlsx [--sheet NAME] [--values] [--tolerance 1e-9] [--json] [-o diff.md]
--values also compares cached formula results (run both through xlsx_recalc.py --write first
if they were saved by scripts).
"""
import argparse
import json
import math
import os
import sys

sys.dont_write_bytecode = True
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import _xlsx_common  # noqa: E402,F401

from openpyxl import load_workbook  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402


def same(a, b, tol):
    if isinstance(a, (int, float)) and isinstance(b, (int, float)) and not isinstance(a, bool):
        return math.isclose(a, b, rel_tol=tol, abs_tol=tol)
    return a == b


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("old")
    ap.add_argument("new")
    ap.add_argument("--sheet")
    ap.add_argument("--values", action="store_true")
    ap.add_argument("--tolerance", type=float, default=1e-9)
    ap.add_argument("--json", action="store_true")
    ap.add_argument("-o", "--output")
    a = ap.parse_args()

    fo, fn = load_workbook(a.old), load_workbook(a.new)
    vo = load_workbook(a.old, data_only=True) if a.values else None
    vn = load_workbook(a.new, data_only=True) if a.values else None
    sheets_o, sheets_n = fo.sheetnames, fn.sheetnames
    result = {"added_sheets": [s for s in sheets_n if s not in sheets_o],
              "removed_sheets": [s for s in sheets_o if s not in sheets_n], "sheets": {}}
    common = [s for s in sheets_o if s in sheets_n and (not a.sheet or s == a.sheet)]
    for s in common:
        wo, wn = fo[s], fn[s]
        changes = []
        rows = max(wo.max_row, wn.max_row)
        cols = max(wo.max_column, wn.max_column)
        for r in range(1, rows + 1):
            for c in range(1, cols + 1):
                x, y = wo.cell(r, c).value, wn.cell(r, c).value
                coord = f"{get_column_letter(c)}{r}"
                if not same(x, y, a.tolerance):
                    kind = "formula" if any(isinstance(v, str) and v.startswith("=") for v in (x, y)) else "value"
                    changes.append({"cell": coord, "kind": kind, "old": x, "new": y})
                elif a.values and isinstance(x, str) and x.startswith("="):
                    rx, ry = vo[s].cell(r, c).value, vn[s].cell(r, c).value
                    if not same(rx, ry, a.tolerance):
                        changes.append({"cell": coord, "kind": "result", "old": rx, "new": ry, "formula": x})
        result["sheets"][s] = {"size_old": f"{wo.max_row}×{wo.max_column}", "size_new": f"{wn.max_row}×{wn.max_column}",
                               "changes": changes}

    if a.json:
        text = json.dumps(result, indent=2, default=str, ensure_ascii=False)
    else:
        lines = [f"# {os.path.basename(a.old)} → {os.path.basename(a.new)}", ""]
        if result["added_sheets"]:
            lines.append(f"Added sheets: {result['added_sheets']}")
        if result["removed_sheets"]:
            lines.append(f"Removed sheets: {result['removed_sheets']}")
        for s, d in result["sheets"].items():
            ch = d["changes"]
            size = "" if d["size_old"] == d["size_new"] else f" (size {d['size_old']} → {d['size_new']})"
            lines.append(f"\n## {s}: {len(ch)} change(s){size}")
            for c in ch[:500]:
                extra = f"   [{c['formula']}]" if c.get("formula") else ""
                lines.append(f"- {c['cell']} ({c['kind']}): {c['old']!r} → {c['new']!r}{extra}")
            if len(ch) > 500:
                lines.append(f"- … {len(ch) - 500} more")
        text = "\n".join(lines)
    if a.output:
        with open(a.output, "w", encoding="utf-8") as fh:
            fh.write(text + "\n")
        print(f"Wrote {a.output}", file=sys.stderr)
    else:
        print(text)


if __name__ == "__main__":
    main()
