"""Shared helpers for xlsx-workbench scripts (imported from the same folder)."""
import os
import shutil
import subprocess
import sys
import tempfile
import warnings

sys.dont_write_bytecode = True
# openpyxl warns about extensions it drops on read (e.g. in LibreOffice-written copies); not useful here
warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

try:
    import openpyxl  # noqa: F401
except ImportError:
    sys.exit("openpyxl is required: pip install openpyxl")

ERROR_VALUES = ("#DIV/0!", "#N/A", "#NAME?", "#NULL!", "#NUM!", "#REF!", "#VALUE!", "#SPILL!", "#CALC!", "#GETTING_DATA")

# A LibreOffice profile setting that forces formulas to be recalculated when a file is
# opened. Without it LibreOffice keeps whatever values Excel cached — possibly stale.
_RECALC_XCU = """<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry" xmlns:xs="http://www.w3.org/2001/XMLSchema" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop></item>
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="ODFRecalcMode" oor:op="fuse"><value>0</value></prop></item>
</oor:items>
"""


def soffice():
    for name in ("soffice", "libreoffice", "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        found = shutil.which(name) or (name if os.path.exists(name) else None)
        if found:
            return found
    return None


def lo_convert(src, fmt, outdir, recalc=True, timeout=300):
    """Convert with LibreOffice (headless, private profile). Returns the output path."""
    exe = soffice()
    if not exe:
        sys.exit("LibreOffice not found (apt install libreoffice-calc / brew install --cask libreoffice)")
    os.makedirs(outdir, exist_ok=True)
    with tempfile.TemporaryDirectory() as profile:
        if recalc:
            os.makedirs(os.path.join(profile, "user"), exist_ok=True)
            with open(os.path.join(profile, "user", "registrymodifications.xcu"), "w") as fh:
                fh.write(_RECALC_XCU)
        cmd = [exe, f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", fmt, "--outdir", outdir, src]
        res = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout)
    out = os.path.join(outdir, os.path.splitext(os.path.basename(src))[0] + "." + fmt.split(":")[0])
    if not os.path.exists(out):
        sys.exit(f"LibreOffice conversion failed: {res.stderr.strip() or res.stdout.strip()}")
    return out


def recalculated_copy(path, tmpdir):
    """Path to a copy of the workbook with every formula recalculated by LibreOffice."""
    ext = os.path.splitext(path)[1].lower().lstrip(".")
    fmt = "xlsx" if ext in ("xlsx", "xlsm", "xls", "ods", "csv") else ext
    return lo_convert(os.path.abspath(path), fmt, tmpdir, recalc=True)


def pick_sheets(wb, spec):
    """spec: None (all), 'Name', '2' (1-based index), or comma list."""
    if not spec:
        return list(wb.worksheets)
    out = []
    for part in str(spec).split(","):
        part = part.strip()
        if part.isdigit() and part not in wb.sheetnames:
            out.append(wb.worksheets[int(part) - 1])
        elif part in wb.sheetnames:
            out.append(wb[part])
        else:
            sys.exit(f"No sheet '{part}'. Sheets: {wb.sheetnames}")
    return out


def default_out(src, suffix, ext=None):
    """Output path next to the input: report.xlsx -> report_<suffix>.<ext>."""
    base, e = os.path.splitext(os.path.abspath(src))
    return f"{base}_{suffix}{'.' + ext if ext else e}"


def inject_cached_values(original, recalculated, out):
    """Write `out` = `original` byte-for-byte, except that every formula cell gets the
    value LibreOffice computed (from `recalculated`) as its cached <v>. All Excel
    features of the original (table styles, charts, validation, VBA parts) are kept,
    and viewers that don't calculate (previews, pandas, data_only=True) see numbers.
    Returns the number of cells updated."""
    import datetime as dt
    import posixpath
    import zipfile
    from lxml import etree
    from openpyxl import load_workbook
    from openpyxl.utils.datetime import to_excel

    NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
          "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
          "pr": "http://schemas.openxmlformats.org/package/2006/relationships"}
    M = "{%s}" % NS["m"]
    vals = load_workbook(recalculated, data_only=True)
    zin = zipfile.ZipFile(original)
    wbx = etree.fromstring(zin.read("xl/workbook.xml"))
    rels = etree.fromstring(zin.read("xl/_rels/workbook.xml.rels"))
    target = {r.get("Id"): r.get("Target") for r in rels.findall("pr:Relationship", NS)}
    sheet_xml = {}
    for sh in wbx.findall("m:sheets/m:sheet", NS):
        t = target.get(sh.get("{%s}id" % NS["r"]))
        if t:
            path = t.lstrip("/") if t.startswith("/") else posixpath.normpath(posixpath.join("xl", t))
            sheet_xml[path] = sh.get("name")
    changed = 0
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            name = sheet_xml.get(item.filename)
            if name and name in vals.sheetnames:
                ws = vals[name]
                root = etree.fromstring(data)
                for c in root.iter(M + "c"):
                    if c.find(M + "f") is None:
                        continue
                    v = ws[c.get("r")].value
                    old = c.find(M + "v")
                    if old is not None:
                        c.remove(old)
                    for attr in ("t",):
                        if attr in c.attrib:
                            del c.attrib[attr]
                    if v is None:
                        continue
                    ve = etree.SubElement(c, M + "v")
                    if isinstance(v, bool):
                        c.set("t", "b")
                        ve.text = "1" if v else "0"
                    elif isinstance(v, (int, float)):
                        ve.text = repr(float(v)) if isinstance(v, float) else str(v)
                    elif isinstance(v, (dt.datetime, dt.date)):
                        ve.text = repr(float(to_excel(v)))
                    elif isinstance(v, str) and v in ERROR_VALUES:
                        c.set("t", "e")
                        ve.text = v
                    else:
                        c.set("t", "str")
                        ve.text = str(v)
                    changed += 1
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            zout.writestr(item, data)
    return changed
