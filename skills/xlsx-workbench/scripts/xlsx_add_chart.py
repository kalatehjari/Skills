#!/usr/bin/env python3
"""Add a native Excel chart (editable in Excel, updates when data changes).

    python xlsx_add_chart.py book.xlsx -o book_chart.xlsx --sheet Lab --type scatter \
        --x "B4:B20" --y "C3:C20,D3:D20" --title "Moisture vs depth" \
        --x-title "Depth (m)" --y-title "Moisture (%)" [--anchor H3] [--target-sheet Charts]
        [--y-reverse] [--size 18x10] [--style 10] [--no-legend] [--marker-only]

--y ranges INCLUDE the header cell (first row = series name); --x excludes it.
--x may be a comma list paired one-to-one with --y (scatter): each series gets its own x.
--series-names "w (%),LL (%)" overrides the names taken from the --y headers.

Depth profile with two parameters (depth down the page, one series per parameter):
    --type scatter --x "B2:B13,C2:C13" --y "A1:A13,A1:A13" --series-names "w (%),LL (%)" \
    --y-reverse --x-title "Water content / LL (%)" --y-title "Depth (m)"
Types: line, scatter (numeric x — use for depth/time/measurements), bar (vertical
columns), barh (horizontal bars), area, pie.
--y-reverse flips the y axis (depth increasing downward, as on borehole logs).
For profiles with depth on the vertical axis, plot depth as Y with --y-reverse.
"""
import argparse
import os
import sys

sys.dont_write_bytecode = True
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import _xlsx_common  # noqa: E402,F401

from openpyxl import load_workbook  # noqa: E402
from openpyxl.chart import AreaChart, BarChart, LineChart, PieChart, Reference, ScatterChart, Series  # noqa: E402
from openpyxl.chart.series import SeriesLabel  # noqa: E402
from openpyxl.chart.shapes import GraphicalProperties  # noqa: E402
from openpyxl.utils.cell import range_boundaries  # noqa: E402


def ref(ws, rng):
    c1, r1, c2, r2 = range_boundaries(rng)
    return Reference(ws, min_col=c1, min_row=r1, max_col=c2, max_row=r2)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx")
    ap.add_argument("-o", "--output", required=True)
    ap.add_argument("--sheet", help="sheet holding the data (default: first)")
    ap.add_argument("--type", choices=["line", "scatter", "bar", "barh", "area", "pie"], default="line")
    ap.add_argument("--x", required=True, help="category / x values, without header, e.g. A4:A20")
    ap.add_argument("--y", required=True, help="comma list of series ranges WITH header, e.g. C3:C20,D3:D20")
    ap.add_argument("--title")
    ap.add_argument("--x-title")
    ap.add_argument("--y-title")
    ap.add_argument("--anchor", help="top-left cell for the chart (default: two columns right of the data)")
    ap.add_argument("--target-sheet", help="put the chart on this sheet (created if missing)")
    ap.add_argument("--size", help="width x height in cm (default 16x9; 12x15 with --y-reverse)")
    ap.add_argument("--style", type=int, default=10)
    ap.add_argument("--y-reverse", action="store_true")
    ap.add_argument("--no-legend", action="store_true")
    ap.add_argument("--marker-only", action="store_true", help="scatter/line: markers without lines")
    ap.add_argument("--series-names", help="comma list of series names (default: header cell of each --y range)")
    ap.add_argument("--x-format", help="number format for the x axis, e.g. 0.0")
    ap.add_argument("--y-format", help="number format for the y axis, e.g. #,##0")
    a = ap.parse_args()
    if os.path.abspath(a.output) == os.path.abspath(a.xlsx):
        sys.exit("Refusing to overwrite the input; choose a different -o")

    keep_vba = a.xlsx.lower().endswith(".xlsm")
    wb = load_workbook(a.xlsx, keep_vba=keep_vba)
    ws = wb[a.sheet] if a.sheet else wb.worksheets[0]
    xranges = [r.strip() for r in a.x.split(",") if r.strip()]
    yranges = [r.strip() for r in a.y.split(",") if r.strip()]
    for xr in xranges:
        c1, r1, c2, r2 = range_boundaries(xr)
        if c1 != c2 and r1 != r2:
            sys.exit(f"--x range {xr} spans several columns and rows; give one column (or a comma list)")
    if len(xranges) not in (1, len(yranges)):
        sys.exit("--x must be one range, or one range per --y range")
    if len(xranges) > 1 and a.type != "scatter":
        sys.exit("Per-series x ranges only work with --type scatter")
    names = [n.strip() for n in a.series_names.split(",")] if a.series_names else []
    palette = ["4472C4", "C0504D", "9BBB59", "8064A2", "F79646", "4BACC6", "7F7F7F", "1F3864"]

    def style_series(s, k):
        color = palette[k % len(palette)]
        s.marker.symbol = ["circle", "square", "triangle", "diamond", "x", "star"][k % 6]
        s.marker.size = 6
        # explicit marker colours: without them LibreOffice draws invisible markers
        s.marker.graphicalProperties = GraphicalProperties(solidFill=color)
        s.marker.graphicalProperties.line.solidFill = color
        if a.marker_only:
            s.graphicalProperties.line.noFill = True
        else:
            s.graphicalProperties.line.solidFill = color
            s.graphicalProperties.line.width = 19050  # 1.5 pt

    if a.type == "scatter":
        ch = ScatterChart()
        ch.scatterStyle = "lineMarker"
        for k, yr in enumerate(yranges):
            xref = ref(ws, xranges[k] if len(xranges) > 1 else xranges[0])
            s = Series(ref(ws, yr), xref, title_from_data=True)
            s.smooth = False
            style_series(s, k)
            ch.series.append(s)
    else:
        xref = ref(ws, xranges[0])
        ch = {"line": LineChart, "bar": BarChart, "barh": BarChart, "area": AreaChart, "pie": PieChart}[a.type]()
        if a.type == "barh":
            ch.type = "bar"
        elif a.type == "bar":
            ch.type = "col"
        for yr in yranges:
            ch.add_data(ref(ws, yr), titles_from_data=True)
        ch.set_categories(xref)
        if a.type == "line":
            for k, s in enumerate(ch.series):
                style_series(s, k)
    for k, s in enumerate(ch.series):
        if k < len(names) and names[k]:
            s.tx = SeriesLabel(v=names[k])
    if a.title:
        ch.title = a.title
    if a.type != "pie":
        if a.x_title:
            ch.x_axis.title = a.x_title
        if a.y_title:
            ch.y_axis.title = a.y_title
        # openpyxl ≥3.1 hides axes unless told otherwise
        ch.x_axis.delete = False
        ch.y_axis.delete = False
        if a.y_reverse:
            ch.y_axis.scaling.orientation = "maxMin"
            if a.type == "scatter":
                ch.x_axis.crosses = "max"  # keep the x axis at the bottom; on a reversed y it would jump to the top
        if a.x_format:
            ch.x_axis.number_format = a.x_format
            ch.x_axis.numFmt.sourceLinked = False if ch.x_axis.numFmt is not None else None
        if a.y_format:
            ch.y_axis.number_format = a.y_format
            ch.y_axis.numFmt.sourceLinked = False if ch.y_axis.numFmt is not None else None
    ch.style = a.style
    size = a.size or ("12x15" if a.y_reverse else "16x9")
    w, h = (float(v) for v in size.lower().split("x"))
    ch.width, ch.height = w, h
    if a.no_legend:
        ch.legend = None

    target = ws
    if a.target_sheet:
        target = wb[a.target_sheet] if a.target_sheet in wb.sheetnames else wb.create_sheet(a.target_sheet)
    anchor = a.anchor or (f"{chr(ord('A') + min(ws.max_column + 1, 24))}2" if target is ws else "A1")
    target.add_chart(ch, anchor)
    wb.save(a.output)
    print(f"Wrote {a.output}: {a.type} chart with {len(yranges)} series "
          f"({', '.join(str(s.tx.v) if s.tx is not None and s.tx.v else (s.tx.strRef.f if s.tx is not None and s.tx.strRef is not None else '?') for s in ch.series)}) "
          f"at {target.title}!{anchor}", file=sys.stderr)
    print("Check it renders: python xlsx_convert.py " + a.output + " --to png", file=sys.stderr)


if __name__ == "__main__":
    main()
