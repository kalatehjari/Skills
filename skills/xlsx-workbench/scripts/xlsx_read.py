#!/usr/bin/env python3
"""Read a sheet (or range) into CSV, JSON, Markdown or a quick summary.

Detects the header row (or use --header-row N; 0 = no header), skips title rows
above it, stops at the first fully blank row unless --all-rows, and fills merged
header cells. Values are the cached results of formulas (what Excel last showed);
add --recalc to compute fresh values with LibreOffice first, or --formulas to get
the formulas themselves.

Usage:
    python xlsx_read.py book.xlsx [--sheet NAME|N] [--range A3:F20] [--header-row 3]
                        [--format csv|json|md|describe] [-o out] [--recalc] [--formulas] [--all-rows]
"describe" prints per-column type, count, min/max/mean/median/sample stdev for
numbers — a fast way to understand a dataset before analysing it.
Formats: csv, json, md (Markdown table, the default), describe.
"""
import argparse
import csv
import datetime as dt
import io
import json
import os
import statistics
import sys
import tempfile

sys.dont_write_bytecode = True
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from _xlsx_common import pick_sheets, recalculated_copy  # noqa: E402

from openpyxl import load_workbook  # noqa: E402


def guess_header(rows):
    for i, r in enumerate(rows[:10]):
        vals = [v for v in r if v not in (None, "")]
        if len({str(v) for v in vals}) >= 2 and sum(isinstance(v, str) for v in vals) / len(vals) >= 0.8:
            return i
    return 0


def fmt(v):
    if isinstance(v, dt.datetime):
        return v.date().isoformat() if (v.hour, v.minute, v.second) == (0, 0, 0) else v.isoformat(sep=" ")
    if isinstance(v, (dt.date, dt.time)):
        return v.isoformat()
    return v


def describe(header, data):
    out = []
    for j, name in enumerate(header):
        col = [r[j] for r in data if j < len(r) and r[j] not in (None, "")]
        nums = [v for v in col if isinstance(v, (int, float)) and not isinstance(v, bool)]
        line = {"column": name, "non_empty": len(col), "missing": len(data) - len(col)}
        if nums and len(nums) >= 0.8 * len(col):
            r = lambda x: round(x, 4)  # noqa: E731
            line.update(type="number", min=r(min(nums)), max=r(max(nums)), mean=r(statistics.fmean(nums)),
                        median=r(statistics.median(nums)))
            if len(nums) > 1:
                line["stdev_sample"] = r(statistics.stdev(nums))
        elif col and all(isinstance(v, (dt.date, dt.datetime)) for v in col):
            line.update(type="date", min=fmt(min(col)), max=fmt(max(col)))
        else:
            distinct = sorted({str(v) for v in col})
            line.update(type="text", distinct=len(distinct), examples=distinct[:5])
        out.append(line)
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx")
    ap.add_argument("--sheet")
    ap.add_argument("--range")
    ap.add_argument("--header-row", type=int, help="1-based row number of the header within the sheet/range; 0 = none")
    ap.add_argument("--format", choices=["csv", "json", "md", "describe"], default="md")
    ap.add_argument("-o", "--output")
    ap.add_argument("--recalc", action="store_true")
    ap.add_argument("--formulas", action="store_true")
    ap.add_argument("--all-rows", action="store_true", help="don't stop at the first blank row")
    a = ap.parse_args()

    with tempfile.TemporaryDirectory() as tmp:
        path = recalculated_copy(a.xlsx, tmp) if a.recalc else a.xlsx
        wb = load_workbook(path, data_only=not a.formulas)
        ws = pick_sheets(wb, a.sheet)[0]
        merged = {}
        for mr in ws.merged_cells.ranges:  # fill merged blocks with their top-left value
            v = ws.cell(mr.min_row, mr.min_col).value
            for r in range(mr.min_row, mr.max_row + 1):
                for c in range(mr.min_col, mr.max_col + 1):
                    merged[(r, c)] = v
        grid = ws[a.range] if a.range else list(ws.iter_rows())
        raw = [[c.value for c in r] for r in grid]
        rows = [[merged.get((c.row, c.column), c.value) for c in r] for r in grid]

    if a.header_row is None:
        h = guess_header(raw)  # on raw cells, so a merged title row isn't mistaken for a header
    else:
        h = a.header_row - 1 if a.header_row > 0 else None
    if h is not None:
        header, seen = [], {}
        for j, v in enumerate(rows[h]):
            name = str(v).strip() if v is not None else f"col{j + 1}"
            seen[name] = seen.get(name, 0) + 1
            header.append(name if seen[name] == 1 else f"{name}_{seen[name]}")  # merged group headers repeat
        body = rows[h + 1:]
    else:
        header = [f"col{j + 1}" for j in range(max(len(r) for r in rows))]
        body = rows
    data, cut_at = [], None
    for k, r in enumerate(body):
        if all(v in (None, "") for v in r):
            if a.all_rows:
                continue
            if data:
                cut_at = k
                break
            continue
        data.append(r)
    if cut_at is not None:
        rest = [r for r in body[cut_at:] if any(v not in (None, "") for v in r)]
        if rest:
            first_row = (h + 2 if h is not None else 1) + cut_at
            print(f"Note: stopped at the blank row {first_row}; {len(rest)} non-empty row(s) below it were not read "
                  f"(often totals/means or a second table). Use --all-rows or --range to include them.", file=sys.stderr)
    # drop trailing columns that are empty everywhere
    width = max([j + 1 for j, n in enumerate(header) if not n.startswith("col")] +
                [max((j + 1 for j, v in enumerate(r) if v not in (None, "")), default=0) for r in data] + [1])
    header = header[:width]
    data = [[fmt(v) for v in r[:width]] for r in data]

    if a.format == "describe":
        text = json.dumps(describe(header, data), indent=2, default=str, ensure_ascii=False) + "\n"
        text = f"Sheet '{ws.title}': {len(data)} data rows, header row {h + 1 if h is not None else '-'}\n" + text
    elif a.format == "json":
        text = json.dumps([dict(zip(header, r)) for r in data], indent=2, default=str, ensure_ascii=False) + "\n"
    elif a.format == "csv":
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(header)
        w.writerows(data)
        text = buf.getvalue()
    else:
        esc = lambda v: "" if v is None else str(v).replace("|", "\\|").replace("\n", " ")  # noqa: E731
        lines = ["| " + " | ".join(esc(x) for x in header) + " |", "|" + "---|" * len(header)]
        lines += ["| " + " | ".join(esc(x) for x in r) + " |" for r in data]
        text = "\n".join(lines) + "\n"
    if a.output:
        with open(a.output, "w", encoding="utf-8", newline="") as fh:
            fh.write(text)
        print(f"Wrote {a.output}: {len(data)} rows × {len(header)} columns from '{ws.title}'", file=sys.stderr)
    else:
        sys.stdout.write(text)


if __name__ == "__main__":
    main()
